import os
import secrets
import json
import re
import csv
import io
import smtplib
import ssl
import uuid
from datetime import date, datetime, timezone
from email.message import EmailMessage
from functools import wraps
from pathlib import Path
from urllib.request import Request as UrlRequest, urlopen

from flask import Flask, flash, redirect, render_template, request, send_file, session, url_for
from flask_sqlalchemy import SQLAlchemy
from flask_socketio import SocketIO, emit, join_room, leave_room
import jwt
from sqlalchemy import func, inspect, text
from werkzeug.security import check_password_hash, generate_password_hash


def load_environment_file():
    env_path = Path(__file__).resolve().parent / ".env"
    if not env_path.exists():
        return
    for raw_line in env_path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip()
        if len(value) >= 2 and value[0] == value[-1] and value[0] in {'"', "'"}:
            value = value[1:-1]
        os.environ.setdefault(key, value)


load_environment_file()

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "dev-only-change-this-secret")
project_database = Path(__file__).resolve().parent / "TaskHQ.db"
default_database = project_database if project_database.exists() else Path(__file__).resolve().parent / "instance" / "TaskHQ.db"
database_url = os.environ.get("DATABASE_URL")
if database_url and database_url.startswith("postgres://"):
    database_url = database_url.replace("postgres://", "postgresql://", 1)
app.config["SQLALCHEMY_DATABASE_URI"] = database_url or f"sqlite:///{default_database}"
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["SQLALCHEMY_ENGINE_OPTIONS"] = {"pool_pre_ping": True}
db = SQLAlchemy(app)
socketio = SocketIO(app, cors_allowed_origins=None, async_mode="threading", manage_session=False)
realtime_connections = {}


def verify_password(stored_password, raw_password):
    if not stored_password or not raw_password:
        return False
    if stored_password == raw_password:
        return True
    if stored_password.startswith('scrypt:') or stored_password.startswith('pbkdf2:') or stored_password.startswith('sha256:') or stored_password.startswith('argon2'):
        try:
            return check_password_hash(stored_password, raw_password)
        except ValueError:
            return False
    try:
        return check_password_hash(stored_password, raw_password)
    except ValueError:
        return False


VALID_ROLES = {"User", "Manager", "Admin"}
VALID_STATUSES = {"Pending", "In Progress", "Completed", "Rejected"}
VALID_PRIORITIES = {"Low", "Medium", "High"}
VALID_DEPENDENCY_TYPES = {"Blocks", "Blocked By", "Related To"}
PERMISSION_SCOPES = {"OWN", "TEAM", "ANY"}
PERMISSION_CATALOG = {
    "Tasks": {
        "tasks.view": "View tasks", "tasks.create": "Create tasks", "tasks.edit": "Edit tasks",
        "tasks.delete": "Delete tasks", "tasks.assign": "Assign tasks", "tasks.reassign": "Reassign tasks",
        "tasks.change_status": "Change task status", "tasks.change_priority": "Change task priority",
        "tasks.change_due_date": "Change due date", "tasks.cancel": "Cancel tasks", "tasks.restore": "Restore tasks",
        "tasks.archive": "Archive tasks", "tasks.export": "Export tasks", "tasks.edit_own": "Edit own tasks",
        "tasks.delete_own": "Delete own tasks", "tasks.assign_own": "Assign own tasks",
    },
    "Users": {"users.view": "View users", "users.create": "Create users", "users.edit": "Edit users", "users.delete": "Delete users", "users.activate": "Activate users", "users.deactivate": "Deactivate users", "users.reset_password": "Reset passwords"},
    "Roles": {"roles.view": "View roles", "roles.create": "Create roles", "roles.edit": "Edit roles", "roles.delete": "Delete roles", "roles.assign": "Assign roles"},
    "Permissions": {"permissions.view": "View permissions", "permissions.manage": "Manage permissions"},
    "Comments": {"comments.view": "View comments", "comments.create": "Create comments", "comments.edit": "Edit comments", "comments.delete": "Delete comments"},
    "Attachments": {"attachments.view": "View attachments", "attachments.upload": "Upload attachments", "attachments.delete": "Delete attachments"},
    "Notifications": {"notifications.view": "View notifications", "notifications.manage": "Manage notifications"},
    "Reports": {"reports.view": "View reports", "reports.create": "Create reports", "reports.export": "Export reports"},
    "Audit Logs": {"audit_logs.view": "View audit logs", "audit_logs.export": "Export audit logs"},
    "Settings": {"settings.view": "View settings", "settings.edit": "Edit settings"},
}
ALL_PERMISSIONS = {permission for group in PERMISSION_CATALOG.values() for permission in group}
TEAM_PERMISSION_OPTIONS = (
    ("view_tasks", "View team tasks"),
    ("create_tasks", "Create and assign tasks"),
    ("post_messages", "Post team messages"),
        ("manage_members", "Manage team members"),
    ("manage_meetings", "Manage live meetings"),
)



class User(db.Model):
    __tablename__ = "users"
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    first_name = db.Column(db.String(50), nullable=False)
    last_name = db.Column(db.String(50), nullable=False)
    username = db.Column(db.String(50), unique=True, nullable=False)
    phone = db.Column(db.String(20))
    email = db.Column(db.String(255))
    gender = db.Column(db.String(10))
    password = db.Column("password_hash" if database_url and database_url.startswith("mssql") else "password", db.String(255), nullable=False)
    role = db.Column(db.String(20), nullable=False, default="User")
    created_at = db.Column(db.String(30))

    @property
    def password_hash(self):
        return self.password

    @password_hash.setter
    def password_hash(self, value):
        self.password = value


class Role(db.Model):
    __tablename__ = "roles"
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    name = db.Column(db.String(80), unique=True, nullable=False)
    description = db.Column(db.String(255))
    is_system = db.Column(db.Boolean, nullable=False, default=False)
    created_at = db.Column(db.String(30), nullable=False)


class Permission(db.Model):
    __tablename__ = "permissions"
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    key = db.Column(db.String(100), unique=True, nullable=False)
    group_name = db.Column(db.String(80), nullable=False)
    label = db.Column(db.String(120), nullable=False)


class RolePermission(db.Model):
    __tablename__ = "role_permissions"
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    role_id = db.Column(db.Integer, db.ForeignKey("roles.id"), nullable=False)
    permission_id = db.Column(db.Integer, db.ForeignKey("permissions.id"), nullable=False)
    scope = db.Column(db.String(10), nullable=False, default="ANY")
    __table_args__ = (db.UniqueConstraint("role_id", "permission_id", name="uq_role_permission"),)


class UserRole(db.Model):
    __tablename__ = "user_roles"
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    user_id = db.Column(db.Integer, db.ForeignKey("users.id"), nullable=False)
    role_id = db.Column(db.Integer, db.ForeignKey("roles.id"), nullable=False)
    assigned_by = db.Column(db.Integer, db.ForeignKey("users.id"))
    created_at = db.Column(db.String(30), nullable=False)
    __table_args__ = (db.UniqueConstraint("user_id", "role_id", name="uq_user_role"),)


class Task(db.Model):
    __tablename__ = "tasks"
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    title = db.Column(db.String(100), nullable=False)
    status = db.Column(db.String(20), nullable=False, default="Pending")
    priority = db.Column(db.String(20), nullable=False, default="Medium")
    due_date = db.Column(db.String(30))
    user_id = db.Column(db.Integer, db.ForeignKey("users.id"))
    creator_id = db.Column(db.Integer, db.ForeignKey("users.id"))
    team_id = db.Column(db.Integer, db.ForeignKey("teams.id"))
    completed_at = db.Column(db.String(30))


class Detail(db.Model):
    __tablename__ = "details"
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    task_id = db.Column(db.Integer, db.ForeignKey("tasks.id"), unique=True)
    description = db.Column(db.Text)
    attachments = db.Column(db.String(255))
    actual_hours = db.Column(db.Float, default=0.0)
    updated_at = db.Column(db.String(30))


class Notification(db.Model):
    __tablename__ = "notifications"
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    user_id = db.Column(db.Integer, db.ForeignKey("users.id"), nullable=False)
    message = db.Column(db.String(255), nullable=False)
    is_read = db.Column(db.Boolean, default=False)
    created_at = db.Column(db.String(30))


class NotificationPreference(db.Model):
    __tablename__ = "notification_preferences"
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    user_id = db.Column(db.Integer, db.ForeignKey("users.id"), unique=True, nullable=False)
    email_enabled = db.Column(db.Boolean, nullable=False, default=True)
    task_created = db.Column(db.Boolean, nullable=False, default=True)
    task_assigned = db.Column(db.Boolean, nullable=False, default=True)
    status_changed = db.Column(db.Boolean, nullable=False, default=True)
    due_reminders = db.Column(db.Boolean, nullable=False, default=True)


class EmailLog(db.Model):
    __tablename__ = "email_logs"
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    user_id = db.Column(db.Integer, db.ForeignKey("users.id"))
    recipient = db.Column(db.String(255), nullable=False)
    subject = db.Column(db.String(255), nullable=False)
    body = db.Column(db.Text, nullable=False)
    status = db.Column(db.String(30), nullable=False, default="mocked")
    error = db.Column(db.Text)
    created_at = db.Column(db.String(30), nullable=False)


class AIRequest(db.Model):
    __tablename__ = "ai_requests"
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    user_id = db.Column(db.Integer, db.ForeignKey("users.id"), nullable=False)
    prompt = db.Column(db.Text, nullable=False)
    response = db.Column(db.Text, nullable=False)
    mode = db.Column(db.String(20), nullable=False, default="mock")
    created_at = db.Column(db.String(30), nullable=False)


class TaskDependency(db.Model):
    __tablename__ = "task_dependencies"
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    predecessor_id = db.Column(db.Integer, db.ForeignKey("tasks.id"), nullable=False)
    successor_id = db.Column(db.Integer, db.ForeignKey("tasks.id"), nullable=False)
    dependency_type = db.Column(db.String(20), nullable=False, default="Blocks")
    created_by = db.Column(db.Integer, db.ForeignKey("users.id"), nullable=False)
    created_at = db.Column(db.String(30), nullable=False)

    __table_args__ = (db.UniqueConstraint("predecessor_id", "successor_id", "dependency_type", name="uq_task_dependency"),)


class Team(db.Model):
    __tablename__ = "teams"
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    name = db.Column(db.String(100), unique=True, nullable=False)
    description = db.Column(db.Text)
    leader_id = db.Column(db.Integer, db.ForeignKey("users.id"))
    status = db.Column(db.String(20), nullable=False, default="Active")
    meeting_url = db.Column(db.String(500))
    meeting_status = db.Column(db.String(30), nullable=False, default="Scheduled")
    meeting_time = db.Column(db.String(30))
    total_meetings_count = db.Column(db.Integer, default=0)
    total_meeting_minutes = db.Column(db.Integer, default=0)
    created_at = db.Column(db.String(30), nullable=False)


class TeamMeeting(db.Model):
    __tablename__ = "team_meetings"
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    team_id = db.Column(db.Integer, db.ForeignKey("teams.id"), nullable=False)
    started_at = db.Column(db.String(30), nullable=False)
    ended_at = db.Column(db.String(30))
    duration_minutes = db.Column(db.Integer, default=0)
    started_by = db.Column(db.Integer, db.ForeignKey("users.id"))



class TeamMember(db.Model):
    __tablename__ = "team_members"
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    team_id = db.Column(db.Integer, db.ForeignKey("teams.id"), nullable=False)
    user_id = db.Column(db.Integer, db.ForeignKey("users.id"), nullable=False)
    permissions = db.Column(db.String(255), nullable=False, default="view_tasks,create_tasks,post_messages")
    created_at = db.Column(db.String(30), nullable=False)
    __table_args__ = (db.UniqueConstraint("team_id", "user_id", name="uq_team_member"),)


class TeamMessage(db.Model):
    __tablename__ = "team_messages"
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    team_id = db.Column(db.Integer, db.ForeignKey("teams.id"), nullable=False)
    user_id = db.Column(db.Integer, db.ForeignKey("users.id"), nullable=False)
    body = db.Column(db.Text, nullable=False)
    created_at = db.Column(db.String(30), nullable=False)


class AuditLog(db.Model):
    __tablename__ = "audit_logs"
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    user_id = db.Column(db.Integer, db.ForeignKey("users.id"))
    action = db.Column(db.String(50), nullable=False)
    entity = db.Column(db.String(50), nullable=False)
    entity_id = db.Column(db.Integer)
    old_value = db.Column(db.Text)
    new_value = db.Column(db.Text)
    ip_address = db.Column(db.String(64))
    created_at = db.Column(db.String(30), nullable=False)


class Comment(db.Model):
    __tablename__ = "comments"
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    task_id = db.Column(db.Integer, db.ForeignKey("tasks.id"), nullable=False)
    user_id = db.Column(db.Integer, db.ForeignKey("users.id"), nullable=False)
    body = db.Column(db.Text, nullable=False)
    created_at = db.Column(db.String(30), nullable=False)


class RealtimeEvent(db.Model):
    __tablename__ = "realtime_events"
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    event_id = db.Column(db.String(64), unique=True, nullable=False, index=True)
    event_type = db.Column(db.String(80), nullable=False, index=True)
    actor_id = db.Column(db.Integer, db.ForeignKey("users.id"))
    entity_type = db.Column(db.String(40), nullable=False)
    entity_id = db.Column(db.Integer)
    payload = db.Column(db.Text, nullable=False, default="{}")
    created_at = db.Column(db.String(30), nullable=False, index=True)


class UserPresence(db.Model):
    __tablename__ = "user_presence"
    user_id = db.Column(db.Integer, db.ForeignKey("users.id"), primary_key=True)
    status = db.Column(db.String(20), nullable=False, default="offline")
    last_seen_at = db.Column(db.String(30), nullable=False)
    updated_at = db.Column(db.String(30), nullable=False)


class AutomationRule(db.Model):
    __tablename__ = "automation_rules"
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    name = db.Column(db.String(120), nullable=False)
    event = db.Column(db.String(40), nullable=False)
    condition = db.Column(db.String(40), nullable=False)
    value = db.Column(db.String(120))
    action = db.Column(db.String(40), nullable=False)
    enabled = db.Column(db.Boolean, nullable=False, default=True)
    created_by = db.Column(db.Integer, db.ForeignKey("users.id"), nullable=False)
    created_at = db.Column(db.String(30), nullable=False)


ROLE_DEFAULTS = {
    "Super Admin": (ALL_PERMISSIONS, "ANY"),
    "Admin": (ALL_PERMISSIONS, "ANY"),
    "Manager": ({
        "tasks.view", "tasks.create", "tasks.edit", "tasks.assign", "tasks.reassign", "tasks.change_status",
        "tasks.change_priority", "tasks.change_due_date", "comments.view", "comments.create", "reports.view",
        "reports.export", "notifications.view", "settings.view", "settings.edit", "users.view",
    }, "TEAM"),
    "User": ({"tasks.view", "tasks.create", "tasks.change_status", "tasks.edit_own", "tasks.assign_own", "comments.view", "comments.create", "notifications.view", "settings.view", "settings.edit"}, "OWN"),
}


def seed_rbac():
    now = datetime.now().isoformat(timespec="seconds")
    permissions = {}
    for group_name, entries in PERMISSION_CATALOG.items():
        for key, label in entries.items():
            permission = Permission.query.filter_by(key=key).first()
            if not permission:
                permission = Permission(key=key, group_name=group_name, label=label)
                db.session.add(permission)
            permissions[key] = permission
    db.session.flush()
    roles = {}
    for name in ROLE_DEFAULTS:
        role = Role.query.filter_by(name=name).first()
        if not role:
            role = Role(name=name, description=f"System role: {name}", is_system=True, created_at=now)
            db.session.add(role)
        roles[name] = role
    db.session.flush()
    for name, (permission_keys, default_scope) in ROLE_DEFAULTS.items():
        role = roles[name]
        existing = {row.permission_id: row for row in RolePermission.query.filter_by(role_id=role.id).all()}
        for key in permission_keys:
            permission = permissions[key]
            row = existing.get(permission.id)
            if not row:
                scope = "TEAM" if name == "User" and key in {"tasks.view", "comments.view"} else default_scope
                db.session.add(RolePermission(role_id=role.id, permission_id=permission.id, scope=scope))
            elif name == "User" and key in {"tasks.view", "comments.view"} and row.scope == "OWN":
                row.scope = "TEAM"
    db.session.flush()
    for user in User.query.all():
        role = roles.get(user.role or "User")
        if role and not UserRole.query.filter_by(user_id=user.id, role_id=role.id).first():
            db.session.add(UserRole(user_id=user.id, role_id=role.id, assigned_by=user.id, created_at=now))
    db.session.commit()


def user_roles(user):
    if isinstance(user, int):
        user = db.session.get(User, user)
    if not user:
        return []
    assigned = list(db.session.query(Role).join(UserRole, UserRole.role_id == Role.id).filter(UserRole.user_id == user.id).all())
    if user.role and not any(role.name == user.role for role in assigned):
        legacy = Role.query.filter_by(name=user.role).first()
        if legacy:
            assigned.append(legacy)
    return assigned


def permission_explanations(user, permission):
    grants = []
    for role in user_roles(user):
        rows = db.session.query(RolePermission, Permission).join(Permission, Permission.id == RolePermission.permission_id).filter(RolePermission.role_id == role.id, Permission.key == permission).all()
        grants.extend({"role": role.name, "scope": row.scope} for row, _permission in rows)
    return grants


def has_permission(user, permission, resource=None, requested_scope=None):
    if isinstance(user, int):
        user = db.session.get(User, user)
    if not user:
        return False
    grants = permission_explanations(user, permission)
    if permission.endswith(".edit") or permission.endswith(".delete") or permission.endswith(".assign"):
        action = permission.rsplit(".", 1)[1]
        if resource is not None:
            owner = getattr(resource, "creator_id", None) == user.id or getattr(resource, "user_id", None) == user.id
            team = bool(getattr(resource, "team_id", None) and TeamMember.query.filter_by(team_id=resource.team_id, user_id=user.id).first())
            if owner:
                grants.extend(permission_explanations(user, f"{permission}_own"))
            if not owner and not team:
                grants = [grant for grant in grants if grant["scope"] == "ANY"]
        grants.extend(permission_explanations(user, f"{permission}_{action}"))
    if not grants:
        return False
    if requested_scope:
        return any(grant["scope"] == requested_scope or grant["scope"] == "ANY" for grant in grants)
    if resource is None:
        return True
    if any(grant["scope"] == "ANY" for grant in grants):
        return True
    owner = getattr(resource, "creator_id", None) == user.id or getattr(resource, "user_id", None) == user.id
    team = bool(getattr(resource, "team_id", None) and TeamMember.query.filter_by(team_id=resource.team_id, user_id=user.id).first())
    return any(grant["scope"] == "OWN" and owner or grant["scope"] == "TEAM" and team for grant in grants)


def can(permission, resource=None):
    return has_permission(session.get("user_id"), permission, resource)


def is_super_admin(user_id):
    return any(role.name == "Super Admin" for role in user_roles(user_id))


def sync_user_roles(user, role_ids, actor_id):
    if not has_permission(actor_id, "roles.assign"):
        return False, "You do not have permission to assign roles."
    selected_roles = Role.query.filter(Role.id.in_(role_ids)).all() if role_ids else []
    if not selected_roles:
        return False, "Select at least one role."
    if any(role.name == "Super Admin" for role in selected_roles) and not is_super_admin(actor_id):
        return False, "Only a Super Admin can assign the Super Admin role."
    UserRole.query.filter_by(user_id=user.id).delete()
    now = datetime.now().isoformat(timespec="seconds")
    for role in selected_roles:
        db.session.add(UserRole(user_id=user.id, role_id=role.id, assigned_by=actor_id, created_at=now))
    legacy_role = next((role.name for role in selected_roles if role.name in VALID_ROLES), "User")
    user.role = legacy_role
    return True, None


def notify_team_live_meeting(team, meeting):
    recipient_ids = {row.user_id for row in TeamMember.query.filter_by(team_id=team.id).all()}
    if meeting.started_by and db.session.get(User, meeting.started_by):
        recipient_ids.add(meeting.started_by)
    if not recipient_ids:
        return

    message = f"team:{team.id}: Live meeting #{meeting.id} started in {team.name}. Join now: {team.meeting_url}"
    existing_recipient_ids = {
        user_id for (user_id,) in db.session.query(Notification.user_id).filter(
            Notification.user_id.in_(recipient_ids),
            Notification.message == message,
        ).all()
    }
    created_at = datetime.now().isoformat(timespec="seconds")
    for recipient_id in recipient_ids - existing_recipient_ids:
        db.session.add(Notification(user_id=recipient_id, message=message, created_at=created_at))


def reconcile_live_meetings():
    active_meetings = TeamMeeting.query.filter(TeamMeeting.ended_at.is_(None)).all()
    active_team_ids = {meeting.team_id for meeting in active_meetings}
    for team in Team.query.filter_by(meeting_status="Live").all():
        if team.id not in active_team_ids:
            if team.meeting_url:
                meeting = TeamMeeting(
                    team_id=team.id,
                    started_at=datetime.now().isoformat(timespec="seconds"),
                    started_by=team.leader_id,
                )
                db.session.add(meeting)
                db.session.flush()
                active_meetings.append(meeting)
                active_team_ids.add(team.id)
            else:
                team.meeting_status = "Scheduled"
    for meeting in active_meetings:
        team = db.session.get(Team, meeting.team_id)
        if not team:
            continue
        team.meeting_status = "Live"
        if team.meeting_url:
            notify_team_live_meeting(team, meeting)
    db.session.commit()


with app.app_context():
    db.create_all()
    alter_add = "ADD" if db.engine.dialect.name == "mssql" else "ADD COLUMN"
    task_columns = {column["name"] for column in inspect(db.engine).get_columns("tasks")}
    user_columns = {column["name"] for column in inspect(db.engine).get_columns("users")}
    if "creator_id" not in {column["name"] for column in inspect(db.engine).get_columns("tasks")}:
        with db.engine.begin() as connection:
            connection.execute(text(f"ALTER TABLE tasks {alter_add} creator_id INTEGER"))
    if "email" not in user_columns:
        with db.engine.begin() as connection:
            connection.execute(text(f"ALTER TABLE users {alter_add} email VARCHAR(255)"))
    if "team_id" not in task_columns:
        with db.engine.begin() as connection:
            connection.execute(text(f"ALTER TABLE tasks {alter_add} team_id INTEGER"))
    if "completed_at" not in task_columns:
        with db.engine.begin() as connection:
            connection.execute(text(f"ALTER TABLE tasks {alter_add} completed_at VARCHAR(30)"))
    team_columns = {column["name"] for column in inspect(db.engine).get_columns("teams")}
    for team_column, team_type in {
                "meeting_url": "VARCHAR(500)",
        "meeting_status": "VARCHAR(30)",
        "meeting_time": "VARCHAR(30)",
        "total_meetings_count": "INTEGER DEFAULT 0",
        "total_meeting_minutes": "INTEGER DEFAULT 0",
    }.items():

        if team_column not in team_columns:
            with db.engine.begin() as connection:
                connection.execute(text(f"ALTER TABLE teams {alter_add} {team_column} {team_type}"))
    team_member_columns = {column["name"] for column in inspect(db.engine).get_columns("team_members")}
    if "permissions" not in team_member_columns:
        with db.engine.begin() as connection:
            connection.execute(text(f"ALTER TABLE team_members {alter_add} permissions VARCHAR(255) DEFAULT 'view_tasks,create_tasks,post_messages'"))
    for team in Team.query.filter(Team.leader_id.isnot(None)).all():
        if db.session.get(User, team.leader_id) and not TeamMember.query.filter_by(team_id=team.id, user_id=team.leader_id).first():
            db.session.add(TeamMember(
                team_id=team.id,
                user_id=team.leader_id,
                permissions="view_tasks,create_tasks,post_messages,manage_members,manage_meetings",
                created_at=datetime.now().isoformat(timespec="seconds"),
            ))
    db.session.commit()
    if not User.query.first():
        db.session.add(User(
            first_name="System",
            last_name="Administrator",
            username=os.environ.get("ADMIN_USERNAME", "admin"),
            password_hash=generate_password_hash(os.environ.get("ADMIN_PASSWORD", "ChangeMe123!")),
            role="Admin",
            created_at=datetime.now().isoformat(timespec="minutes"),
        ))
        db.session.commit()
    default_accounts = (
        ("admin", "Admin", "Admin", "Administrator"),
        ("user", "User", "Standard", "User"),
        ("manager", "Manager", "Task", "Manager"),
    )
    default_passwords = {
        "admin": os.environ.get("ADMIN_PASSWORD", "Admin@12345"),
        "user": os.environ.get("USER_PASSWORD", "User@12345"),
        "manager": os.environ.get("MANAGER_PASSWORD", "Manager@12345"),
    }
    for username, role, first_name, last_name in default_accounts:
        if not User.query.filter(func.lower(func.trim(User.username)) == username).first():
            db.session.add(User(
                first_name=first_name,
                last_name=last_name,
                username=username,
                password_hash=generate_password_hash(default_passwords[username]),
                role=role,
                created_at=datetime.now().isoformat(timespec="minutes"),
            ))
    db.session.flush()
    seed_rbac()
    reconcile_live_meetings()


def get_csrf_token():
    if "csrf_token" not in session:
        session["csrf_token"] = secrets.token_urlsafe(32)
    return session["csrf_token"]


def valid_csrf():
    submitted = request.form.get("csrf_token", "") or request.headers.get("X-CSRF-Token", "")
    return secrets.compare_digest(session.get("csrf_token", ""), submitted)


def realtime_token(user_id):
    now = int(datetime.now(timezone.utc).timestamp())
    return jwt.encode(
        {"sub": str(user_id), "iat": now, "exp": now + 3600},
        app.secret_key,
        algorithm="HS256",
    )


def realtime_event(event_type, actor_id, entity_type, entity_id=None, payload=None, rooms=None):
    """Persist an event before publishing it so reconnect sync has a durable source."""
    created_at = datetime.now().isoformat(timespec="seconds")
    event = RealtimeEvent(
        event_id=uuid.uuid4().hex,
        event_type=event_type,
        actor_id=actor_id,
        entity_type=entity_type,
        entity_id=entity_id,
        payload=json.dumps(payload or {}, default=str),
        created_at=created_at,
    )
    db.session.add(event)
    db.session.flush()
    message = {
        "sequence": event.id,
        "eventId": event.event_id,
        "eventType": event.event_type,
        "timestamp": event.created_at,
        "userId": event.actor_id,
        "entityType": event.entity_type,
        "entityId": event.entity_id,
        "payload": payload or {},
    }
    return message, rooms or []


def publish_realtime(message, rooms):
    for room in rooms:
        socketio.emit("realtime_event", message, to=room)


def record_realtime_event(event_type, actor_id, entity_type, entity_id=None, payload=None, rooms=None):
    message, target_rooms = realtime_event(event_type, actor_id, entity_type, entity_id, payload, rooms)
    db.session.commit()
    publish_realtime(message, target_rooms)


def can_access_task(user_id, task):
    if not task:
        return False
    user = db.session.get(User, user_id)
    if user and user.role == "Admin":
        return True
    team_member = TeamMember.query.filter_by(team_id=task.team_id, user_id=user_id).first() if task.team_id else None
    return bool(task.user_id == user_id or task.creator_id == user_id or team_member)


@app.route("/api/realtime/token")
def realtime_auth_token():
    if not session.get("user_id"):
        return {"ok": False, "error": "Unauthorized"}, 401
    return {"token": realtime_token(session["user_id"]), "user_id": session["user_id"]}


@socketio.on("connect")
def realtime_connect(auth=None):
    auth = auth or {}
    token = auth.get("token")
    if not token:
        return False
    try:
        claims = jwt.decode(token, app.secret_key, algorithms=["HS256"])
        user_id = int(claims["sub"])
    except (jwt.InvalidTokenError, KeyError, TypeError, ValueError):
        return False
    user = db.session.get(User, user_id)
    if not user:
        return False
    realtime_connections[request.sid] = user_id
    join_room("global")
    join_room(f"user:{user_id}")
    now = datetime.now().isoformat(timespec="seconds")
    presence = db.session.get(UserPresence, user_id) or UserPresence(user_id=user_id, last_seen_at=now, updated_at=now)
    presence.status = "online"
    presence.last_seen_at = now
    presence.updated_at = now
    db.session.add(presence)
    db.session.commit()
    socketio.emit("presence", {"eventType": "USER_ONLINE", "userId": user_id, "status": "online", "timestamp": now}, to="global")
    emit("connected", {"userId": user_id, "serverTime": now})


@socketio.on("heartbeat")
def realtime_heartbeat():
    user_id = realtime_connections.get(request.sid)
    if not user_id:
        return {"ok": False, "error": "Unauthorized"}
    now = datetime.now().isoformat(timespec="seconds")
    presence = db.session.get(UserPresence, user_id)
    if presence:
        presence.status = "online"
        presence.last_seen_at = now
        presence.updated_at = now
        db.session.commit()
    return {"ok": True, "timestamp": now}


@socketio.on("join_room")
def realtime_join_room(data):
    user_id = realtime_connections.get(request.sid)
    room = (data or {}).get("room", "")
    if not user_id or not isinstance(room, str):
        return {"ok": False, "error": "Unauthorized"}
    parts = room.split(":", 1)
    allowed = room == "global"
    if len(parts) == 2 and parts[0] == "user":
        allowed = int(parts[1]) == user_id if parts[1].isdigit() else False
    elif len(parts) == 2 and parts[0] == "task" and parts[1].isdigit():
        allowed = can_access_task(user_id, db.session.get(Task, int(parts[1])))
    elif len(parts) == 2 and parts[0] == "team" and parts[1].isdigit():
        team = db.session.get(Team, int(parts[1]))
        allowed = bool(team and (db.session.get(User, user_id).role in {"Admin", "Manager"} or TeamMember.query.filter_by(team_id=team.id, user_id=user_id).first()))
    if not allowed:
        return {"ok": False, "error": "Room access denied"}
    join_room(room)
    return {"ok": True, "room": room}


@socketio.on("sync")
def realtime_sync(data):
    user_id = realtime_connections.get(request.sid)
    if not user_id:
        return {"ok": False, "error": "Unauthorized"}
    since_id = int((data or {}).get("lastEventId", 0) or 0)
    events = RealtimeEvent.query.filter(RealtimeEvent.id > since_id).order_by(RealtimeEvent.id.asc()).limit(200).all()
    result = []
    for event in events:
        payload = json.loads(event.payload or "{}")
        if event.entity_type == "task" and not can_access_task(user_id, db.session.get(Task, event.entity_id)) and user_id not in payload.get("recipientIds", []):
            continue
        if event.entity_type == "team":
            team = db.session.get(Team, event.entity_id)
            if not team or (db.session.get(User, user_id).role not in {"Admin", "Manager"} and not TeamMember.query.filter_by(team_id=team.id, user_id=user_id).first()):
                continue
        result.append({"sequence": event.id, "eventId": event.event_id, "eventType": event.event_type, "timestamp": event.created_at, "userId": event.actor_id, "entityType": event.entity_type, "entityId": event.entity_id, "payload": payload})
    return {"ok": True, "events": result, "lastEventId": events[-1].id if events else since_id}


@socketio.on("disconnect")
def realtime_disconnect():
    user_id = realtime_connections.pop(request.sid, None)
    if not user_id:
        return
    now = datetime.now().isoformat(timespec="seconds")
    presence = db.session.get(UserPresence, user_id)
    if presence:
        presence.status = "offline"
        presence.last_seen_at = now
        presence.updated_at = now
        db.session.commit()
    socketio.emit("presence", {"eventType": "USER_OFFLINE", "userId": user_id, "status": "offline", "timestamp": now}, to="global")


def write_audit(action, entity, entity_id=None, old_value=None, new_value=None):
    db.session.add(AuditLog(user_id=session.get("user_id"), action=action, entity=entity, entity_id=entity_id,
                            old_value=json.dumps(old_value, default=str) if old_value is not None else None,
                            new_value=json.dumps(new_value, default=str) if new_value is not None else None,
                            ip_address=request.headers.get("X-Forwarded-For", request.remote_addr),
                            created_at=datetime.now().isoformat(timespec="seconds")))


def parse_import_file(upload):
    filename = (upload.filename or "").lower()
    raw = upload.read()
    if filename.endswith(".csv"):
        return list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig")))), "csv"
    if filename.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise ValueError("Excel import requires the openpyxl package.") from error
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.values)
        if not rows:
            return [], "xlsx"
        headers = [str(value or "").strip() for value in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:]], "xlsx"
    raise ValueError("Only CSV and Excel files are supported.")


def validate_import_rows(resource, rows):
    valid, errors = [], []
    required = {"tasks": {"title"}, "users": {"username", "password"}, "teams": {"name", "leader_username"}}[resource]
    for index, row in enumerate(rows, start=2):
        normalized = {str(key).strip().lower(): (str(value).strip() if value is not None else "") for key, value in row.items()}
        missing = sorted(field for field in required if not normalized.get(field))
        if missing:
            errors.append({"row": index, "error": f"Missing: {', '.join(missing)}"})
        else:
            valid.append(normalized)
    return valid, errors


@app.context_processor
def inject_template_data():
    current_user = User.query.get(session["user_id"]) if session.get("user_id") else None
    return {"current_user": current_user, "csrf_token": get_csrf_token(), "can": can}


@app.after_request
def add_security_headers(response):
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "SAMEORIGIN"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    return response


def login_required(view):
    @wraps(view)
    def wrapped_view(*args, **kwargs):
        if not session.get("user_id"):
            flash("Please sign in to continue.", "danger")
            return redirect(url_for("login"))
        return view(*args, **kwargs)
    return wrapped_view


def permission_required(permission, resource=None):
    user_id = session.get("user_id")
    target = resource() if callable(resource) else resource
    if not has_permission(user_id, permission, target):
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return {"ok": False, "error": "You are not authorized for this action.", "permission": permission}, 403
        flash("You do not have permission to perform this action.", "danger")
        return redirect(url_for("tasks"))
    return None


def admin_required(view):
    @wraps(view)
    @login_required
    def wrapped_view(*args, **kwargs):
        if not has_permission(session.get("user_id"), "permissions.manage"):
            flash("Administrator permission is required.", "danger")
            return redirect(url_for("tasks"))
        return view(*args, **kwargs)
    return wrapped_view


def get_shared_data():
    if not session.get("user_id"):
        return [], 0
    notifications = Notification.query.filter_by(user_id=session["user_id"], is_read=False).order_by(Notification.id.desc()).all()
    return notifications, len(notifications)


def get_notification_preferences(user_id):
    preferences = NotificationPreference.query.filter_by(user_id=user_id).first()
    if not preferences:
        preferences = NotificationPreference(user_id=user_id)
        db.session.add(preferences)
        db.session.commit()
    return preferences


def smtp_configured():
    return bool(os.environ.get("SMTP_HOST") and os.environ.get("SMTP_FROM"))


def send_email_notification(recipient, subject, body, user_id=None, preference_name=None):
    if not recipient:
        return "skipped"
    if user_id and not get_notification_preferences(user_id).email_enabled:
        return "disabled"
    if user_id and preference_name and not getattr(get_notification_preferences(user_id), preference_name, True):
        return "disabled"
    log = EmailLog(user_id=user_id, recipient=recipient, subject=subject, body=body,
                   status="mocked" if not smtp_configured() else "pending",
                   created_at=datetime.now().isoformat(timespec="seconds"))
    db.session.add(log)
    db.session.flush()
    if not smtp_configured():
        return "mocked"
    try:
        message = EmailMessage()
        message["From"] = os.environ["SMTP_FROM"]
        message["To"] = recipient
        message["Subject"] = subject
        message.set_content(body)
        host = os.environ["SMTP_HOST"]
        port = int(os.environ.get("SMTP_PORT", "587"))
        username = os.environ.get("SMTP_USERNAME")
        password = os.environ.get("SMTP_PASSWORD")
        context = ssl.create_default_context()
        with smtplib.SMTP(host, port, timeout=15) as smtp:
            if os.environ.get("SMTP_USE_TLS", "true").lower() == "true":
                smtp.starttls(context=context)
            if username and password:
                smtp.login(username, password)
            smtp.send_message(message)
        log.status = "sent"
        return "sent"
    except Exception as error:
        log.status = "failed"
        log.error = str(error)
        return "failed"


def visible_tasks_for_user(user_id):
    query = Task.query
    if session.get("role") != "Admin":
        team_ids = db.session.query(TeamMember.team_id).filter(TeamMember.user_id == user_id)
        query = query.filter((Task.user_id == user_id) | (Task.creator_id == user_id) | Task.team_id.in_(team_ids))
    return query.order_by(Task.id.desc()).all()


def mock_ai_answer(prompt, user_id):
    tasks = visible_tasks_for_user(user_id)
    lowered = prompt.casefold()
    is_arabic = any("\u0600" <= character <= "\u06ff" for character in prompt)
    today = date.today().isoformat()
    week_end = date.fromordinal(date.today().toordinal() + 7).isoformat()
    overdue = [task for task in tasks if task.due_date and task.status != "Completed" and task.due_date[:10] < today]
    due_this_week = [task for task in tasks if task.due_date and today <= task.due_date[:10] <= week_end and task.status != "Completed"]
    completed = [task for task in tasks if task.status == "Completed"]
    active = [task for task in tasks if task.status != "Completed"]

    def task_list(rows):
        if not rows:
            return "لا توجد مهام مطابقة." if is_arabic else "No matching tasks."
        usernames = {
            user.id: user.username
            for user in User.query.filter(User.id.in_({task.user_id for task in rows if task.user_id})).all()
        }
        lines = []
        for task in rows[:12]:
            assignee = usernames.get(task.user_id, "غير معيّن" if is_arabic else "Unassigned")
            due = task.due_date[:10] if task.due_date else ("بلا موعد" if is_arabic else "No due date")
            if is_arabic:
                lines.append(f"{task.title} — الحالة: {task.status}، الأولوية: {task.priority}، الموعد: {due}، المسؤول: {assignee}")
            else:
                lines.append(f"{task.title} — {task.status}, {task.priority} priority, due {due}, assigned to {assignee}")
        return "\n".join(lines)

    if any(word in lowered for word in ("overdue", "late", "متأخر", "متأخرة", "متأخرة")):
        return (f"لديك {len(overdue)} مهمة متأخرة:\n" if is_arabic else f"You have {len(overdue)} overdue task(s):\n") + task_list(overdue)
    if any(word in lowered for word in ("this week", "due this week", "الأسبوع", "هذا الأسبوع")):
        return (f"المهام المستحقة خلال الأيام السبعة القادمة ({len(due_this_week)}):\n" if is_arabic else f"Tasks due in the next seven days ({len(due_this_week)}):\n") + task_list(due_this_week)
    if any(word in lowered for word in ("completed", "done", "مكتمل", "مكتملة", "المنجزة", "المنجز")):
        return (f"المهام المكتملة ({len(completed)}):\n" if is_arabic else f"Completed tasks ({len(completed)}):\n") + task_list(completed)
    if any(word in lowered for word in ("high priority", "urgent", "أولوية عالية", "عاجل", "عاجلة")):
        return task_list([task for task in active if task.priority == "High"])
    if any(word in lowered for word in ("low priority", "أولوية منخفضة")):
        return task_list([task for task in active if task.priority == "Low"])
    if any(word in lowered for word in ("in progress", "قيد التنفيذ", "جارية", "جاري تنفيذ")):
        return task_list([task for task in tasks if task.status == "In Progress"])
    if any(word in lowered for word in ("pending", "not started", "معلقة", "معلق", "قيد الانتظار", "لم تبدأ")):
        return task_list([task for task in tasks if task.status == "Pending"])

    matching_tasks = [task for task in tasks if task.title.casefold() in lowered]
    if matching_tasks:
        return task_list(matching_tasks)

    if any(word in lowered for word in ("subtask", "steps", "خطوات", "خطوة", "قسّم", "قسم")):
        return "الخطوات المقترحة: حدّد النتيجة، قسّم العمل إلى مهام صغيرة، عيّن مسؤولًا لكل مهمة، حدّد موعدًا، ثم راجع الإنجاز." if is_arabic else "Suggested steps: define the outcome, split the work into small tasks, assign an owner to each, set due dates, and review completion."
    if any(word in lowered for word in ("priority", "أولوية", "الأولويات")):
        return "استخدم عالية للعمل العاجل أو المرتبط بموعد قريب، ومتوسطة للعمل المخطط، ومنخفضة للمتابعات المرنة." if is_arabic else "Use High for urgent or deadline-sensitive work, Medium for planned work, and Low for flexible follow-ups."
    if any(word in lowered for word in ("category", "categorize", "تصنيف", "صنّف", "صنف")):
        return "يمكنك تصنيف المهام حسب التشغيل أو التخطيط أو التسليم أو المتابعة." if is_arabic else "Useful categories include Operations, Planning, Delivery, and Follow-up."
    if any(word in lowered for word in ("summar", "overview", "لخص", "تلخيص", "ملخص", "ملخّص", "نظرة عامة")):
        if is_arabic:
            return f"ملخص مهامك الظاهرة: {len(tasks)} إجمالًا، {len(completed)} مكتملة، {len(active)} مفتوحة، و{len(overdue)} متأخرة."
        return f"Your visible tasks: {len(tasks)} total, {len(completed)} completed, {len(active)} open, and {len(overdue)} overdue."
    if any(word in lowered for word in ("task", "tasks", "مهامي", "المهام", "اعرض", "أعرض", "قائمة")):
        return (f"هذه مهامك الظاهرة ({len(tasks)}):\n" if is_arabic else f"Your visible tasks ({len(tasks)}):\n") + task_list(tasks)
    if is_arabic:
        return "أستطيع مساعدتك في عرض مهامك، المتأخر منها، المستحق هذا الأسبوع، المكتمل، حسب الأولوية أو الحالة، أو تلخيص مهامك."
    return "I can list your tasks, find overdue or due-soon work, filter by status or priority, show completed tasks, or summarize your workload."


def ai_answer(prompt, user_id):
    api_key = os.environ.get("AI_API_KEY")
    if not api_key:
        return mock_ai_answer(prompt, user_id), "mock"
    tasks = visible_tasks_for_user(user_id)
    context = [{"title": task.title, "status": task.status, "priority": task.priority, "due_date": task.due_date} for task in tasks]
    payload = {"model": os.environ.get("AI_MODEL", "gpt-4o-mini"), "messages": [
        {"role": "system", "content": "Answer only from the user's visible task data. Never invent or reveal other users' data."},
        {"role": "user", "content": f"Visible tasks: {json.dumps(context)}\nQuestion: {prompt}"},
    ], "temperature": 0.2}
    endpoint = os.environ.get("AI_API_URL", "https://api.openai.com/v1/chat/completions")
    try:
        request = UrlRequest(endpoint, data=json.dumps(payload).encode(), headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"}, method="POST")
        with urlopen(request, timeout=30) as response:
            result = json.loads(response.read().decode())
        return result["choices"][0]["message"]["content"], "provider"
    except Exception:
        return mock_ai_answer(prompt, user_id), "mock-fallback"


def task_rows_for_current_user():
    query = db.session.query(
        Task.id, Task.title, Task.status, Task.priority, Task.due_date,
        User.username, Detail.description, Detail.attachments, Detail.actual_hours, Task.creator_id, Task.user_id,
    ).outerjoin(User, Task.user_id == User.id).outerjoin(Detail, Task.id == Detail.task_id)
    if session.get("role") != "Admin":
        team_ids = db.session.query(TeamMember.team_id).filter(TeamMember.user_id == session["user_id"])
        query = query.filter((Task.user_id == session["user_id"]) | (Task.creator_id == session["user_id"]) | Task.team_id.in_(team_ids))
    return query.order_by(Task.id.desc()).all()


def dependency_predecessors(task_id):
    return db.session.query(Task).join(
        TaskDependency, TaskDependency.predecessor_id == Task.id
    ).filter(
        TaskDependency.successor_id == task_id,
        TaskDependency.dependency_type.in_(["Blocks", "Blocked By"]),
    ).all()


def dependency_would_cycle(predecessor_id, successor_id):
    pending = [predecessor_id]
    visited = set()
    while pending:
        current = pending.pop()
        if current == successor_id:
            return True
        if current in visited:
            continue
        visited.add(current)
        pending.extend(
            row.successor_id for row in TaskDependency.query.filter_by(predecessor_id=current).all()
        )
    return False


def run_automation_once():
    """Evaluate enabled rules once; callable from a scheduler or worker process."""
    fired = 0
    for rule in AutomationRule.query.filter_by(enabled=True).all():
        tasks = Task.query.filter(Task.status == rule.value).all() if rule.event == "status_changed" and rule.condition in {"equals", "changes_to"} else Task.query.filter(Task.due_date < date.today().isoformat(), Task.status != "Completed").all() if rule.event == "overdue" else []
        for task in tasks:
            recipient = db.session.get(User, task.creator_id or task.user_id)
            if not recipient:
                continue
            message = f"Automation '{rule.name}' matched task: {task.title}"
            db.session.add(Notification(user_id=recipient.id, message=message, created_at=datetime.now().isoformat(timespec="seconds")))
            if rule.action == "email":
                send_email_notification(recipient.email, rule.name, message, recipient.id)
            fired += 1
    db.session.commit()
    return fired


@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("user_id"):
        return redirect(url_for("tasks"))
    if request.method == "POST":
        username = request.form.get("login_username", request.form.get("username", "")).strip().lower()
        user = User.query.filter(func.lower(func.trim(User.username)) == username).first()
        password = request.form.get("login_password", request.form.get("password", ""))
        if user and verify_password(user.password_hash, password):
            if user.password_hash == password:
                user.password_hash = generate_password_hash(password)
                db.session.commit()
            session.clear()
            session.update(user_id=user.id, username=user.username, role=user.role or "User")
            get_csrf_token()
            flash("Logged in successfully.", "success")
            return redirect(url_for("tasks"))
        flash("Invalid username or password.", "danger")
    return render_template("login.html")


@app.route("/logout", methods=["POST"])
@login_required
def logout():
    if not valid_csrf():
        flash("Invalid request. Please try again.", "danger")
        return redirect(url_for("tasks"))
    session.clear()
    flash("Logged out successfully.", "success")
    return redirect(url_for("login"))


@app.route("/notifications/read", methods=["POST"])
@login_required
def mark_notifications_read():
    if not valid_csrf():
        return {"ok": False, "error": "Invalid request"}, 400
    Notification.query.filter_by(user_id=session["user_id"], is_read=False).update({"is_read": True})
    db.session.commit()
    return {"ok": True}


@app.route("/notifications/feed")
@login_required
def notifications_feed():
    user_id = session["user_id"]
    notifications = Notification.query.filter_by(user_id=user_id).order_by(Notification.id.desc()).limit(20).all()
    unread_count = Notification.query.filter_by(user_id=user_id, is_read=False).count()
    items = []
    for notification in notifications:
        message = notification.message or ""
        if message.startswith("team:") or message.startswith("team_task:"):
            parts = message.split(":", 2)
            if len(parts) == 3:
                message = parts[2].strip()
        items.append({
            "id": notification.id,
            "message": message,
            "created_at": notification.created_at,
            "url": url_for("open_notification", notification_id=notification.id),
            "is_read": notification.is_read,
        })
    return {"unread_count": unread_count, "notifications": items}


@app.route("/notifications/<int:notification_id>")
@login_required
def open_notification(notification_id):
    notification = Notification.query.filter_by(id=notification_id, user_id=session["user_id"]).first()
    if not notification:
        return redirect(url_for("tasks"))

    notification.is_read = True
    db.session.commit()
    message = notification.message or ""

    if message.startswith("team_task:"):
        task_id = message.split(":", 2)[1]
        if task_id.isdigit() and db.session.get(Task, int(task_id)):
            return redirect(url_for("task_view", task_id=int(task_id)))

    if message.startswith("team:"):
        team_id = message.split(":", 2)[1]
        if team_id.isdigit() and db.session.get(Team, int(team_id)):
            return redirect(url_for("team_detail", team_id=int(team_id)))

    if message.startswith("New team task assigned:") or message.startswith("New task assigned:"):
        title = message.split(":", 1)[1].strip()
        task = Task.query.filter_by(title=title, user_id=session["user_id"]).order_by(Task.id.desc()).first()
        if task:
            return redirect(url_for("task_view", task_id=task.id))

    if " posted in " in message:
        team_name = message.split(" posted in ", 1)[1].split(":", 1)[0].strip()
        team = Team.query.filter_by(name=team_name).first()
        if team:
            return redirect(url_for("team_detail", team_id=team.id))

    if "task comment:" in message:
        title = message.split("task comment:", 1)[1].strip()
        task = Task.query.filter_by(title=title).order_by(Task.id.desc()).first()
        if task:
            return redirect(url_for("task_view", task_id=task.id))

    return redirect(url_for("tasks"))


@app.route("/tasks/<int:task_id>/view")
@login_required
def task_view(task_id):
    task = db.session.get(Task, task_id)
    member_team_ids = {team_id for (team_id,) in db.session.query(TeamMember.team_id).filter_by(user_id=session["user_id"]).all()}
    if not task or not has_permission(session["user_id"], "tasks.view", task) or (not can_access_task(session["user_id"], task) and task.team_id not in member_team_ids):
        flash("You do not have access to this task.", "danger")
        return redirect(url_for("tasks"))
    detail = Detail.query.filter_by(task_id=task.id).first()
    comments = Comment.query.filter_by(task_id=task.id).order_by(Comment.id.asc()).all()
    users = {user.id: user for user in User.query.order_by(User.username).all()}
    return render_template(
        "task_details.html",
        task=task,
        detail=detail,
        comments=comments,
        users=users,
        can_change_status=has_permission(session["user_id"], "tasks.change_status", task),
        can_manage_task=has_permission(session["user_id"], "tasks.edit", task),
        notifications=get_shared_data()[0],
        unread_count=get_shared_data()[1],
    )


@app.route("/assistant", methods=["GET", "POST"])
@login_required
def assistant():
    answer = None
    mode = "provider" if os.environ.get("AI_API_KEY") else "mock"
    prompt = ""
    if request.method == "POST":
        if not valid_csrf():
            flash("Invalid request. Please refresh and try again.", "danger")
            return redirect(url_for("assistant"))
        prompt = request.form.get("prompt", "").strip()
        if prompt:
            answer, mode = ai_answer(prompt, session["user_id"])
            db.session.add(AIRequest(user_id=session["user_id"], prompt=prompt, response=answer, mode=mode, created_at=datetime.now().isoformat(timespec="seconds")))
            db.session.commit()
        else:
            flash("Ask the assistant a question first.", "danger")
    history = AIRequest.query.filter_by(user_id=session["user_id"]).order_by(AIRequest.id.desc()).limit(12).all()
    history.reverse()
    if answer and history and history[-1].response == answer:
        history.pop()
    notifications, unread_count = get_shared_data()
    return render_template("assistant.html", answer=answer, mode=mode, prompt=prompt, history=history, notifications=notifications, unread_count=unread_count)


@app.route("/api/ai/ask", methods=["POST"])
@login_required
def api_ai_ask():
    if not valid_csrf():
        return {"ok": False, "error": "Invalid request"}, 400
    payload = request.get_json(silent=True) or request.form
    prompt = str(payload.get("prompt", "")).strip()
    if not prompt or len(prompt) > 2000:
        return {"ok": False, "error": "Prompt is required"}, 400
    answer, mode = ai_answer(prompt, session["user_id"])
    db.session.add(AIRequest(user_id=session["user_id"], prompt=prompt, response=answer, mode=mode, created_at=datetime.now().isoformat(timespec="seconds")))
    db.session.commit()
    return {"ok": True, "answer": answer, "mode": mode}


@app.route("/settings", methods=["GET", "POST"])
@login_required
def settings():
    if not has_permission(session["user_id"], "settings.view"):
        flash("You do not have permission to view settings.", "danger")
        return redirect(url_for("tasks"))
    preferences = get_notification_preferences(session["user_id"])
    user = db.session.get(User, session["user_id"])
    if request.method == "POST":
        if not has_permission(session["user_id"], "settings.edit"):
            flash("You do not have permission to edit settings.", "danger")
            return redirect(url_for("settings"))
        if not valid_csrf():
            flash("Invalid request. Please try again.", "danger")
            return redirect(url_for("settings"))

        current_password = request.form.get("current_password", "")
        new_password = request.form.get("new_password", "")
        confirm_password = request.form.get("confirm_password", "")

        if current_password or new_password or confirm_password:
            if not current_password or not new_password or not confirm_password:
                flash("Please fill in the current password, new password, and confirmation.", "danger")
                return redirect(url_for("settings"))
            if not verify_password(user.password_hash, current_password):
                flash("Current password is incorrect.", "danger")
                return redirect(url_for("settings"))
            if len(new_password) < 8:
                flash("New password must be at least 8 characters long.", "danger")
                return redirect(url_for("settings"))
            if new_password != confirm_password:
                flash("New password and confirmation do not match.", "danger")
                return redirect(url_for("settings"))
            user.password_hash = generate_password_hash(new_password)
            flash("Password updated successfully.", "success")

        user.email = request.form.get("email", "").strip() or None
        for field in ("email_enabled", "task_created", "task_assigned", "status_changed", "due_reminders"):
            setattr(preferences, field, request.form.get(field) == "on")
        db.session.commit()
        flash("Notification settings saved.", "success")
        return redirect(url_for("settings"))
    notifications, unread_count = get_shared_data()
    return render_template("settings.html", user=user, preferences=preferences,
                           ai_mode="provider" if os.environ.get("AI_API_KEY") else "mock",
                           notifications=notifications, unread_count=unread_count)


@app.route("/settings/test-email", methods=["POST"])
@login_required
def test_email():
    if not valid_csrf():
        flash("Invalid request. Please try again.", "danger")
        return redirect(url_for("settings"))
    user = db.session.get(User, session["user_id"])
    result = send_email_notification(user.email if user else None, "TK test email", "This is a test notification from your task workspace.", user.id if user else None)

    db.session.commit()
    if result == "sent":
        flash("Test email sent successfully.", "success")
    elif result == "mocked":
        flash("Test email recorded in mock mode. Add SMTP settings to send real email.", "success")
    else:
        flash("Add an email address before testing email notifications.", "danger")
    return redirect(url_for("settings"))


@app.route("/")
@login_required
def index():
    rows = task_rows_for_current_user()
    notifications, unread_count = get_shared_data()
    return render_template("index.html", tasks_count=len(rows), users_count=User.query.count(), notifications=notifications, unread_count=unread_count)


@app.route("/teams", methods=["GET", "POST"])
@login_required
def teams():
    if request.method == "POST":
        if session.get("role") not in {"Admin", "Manager"}:
            flash("Team management permission is required.", "danger")
            return redirect(url_for("teams"))
        if not valid_csrf():
            flash("Invalid request. Please try again.", "danger")
            return redirect(url_for("teams"))
        action = request.form.get("action", "create")
        team_id = request.form.get("team_id", type=int)
        if action == "delete":
            team = db.session.get(Team, team_id)
            if team:
                Task.query.filter_by(team_id=team.id).update({"team_id": None})
                TeamMember.query.filter_by(team_id=team.id).delete()
                TeamMeeting.query.filter_by(team_id=team.id).delete()
                TeamMessage.query.filter_by(team_id=team.id).delete()
                db.session.delete(team)
                db.session.commit()
                flash("Team deleted successfully.", "success")
            return redirect(url_for("teams"))
        if action == "member":
            member_id = request.form.get("user_id", type=int)
            team = db.session.get(Team, team_id)
            member = db.session.get(User, member_id)
            if team and member:
                existing_member = TeamMember.query.filter_by(team_id=team.id, user_id=member.id).first()
                if existing_member:
                    flash("This user is already a team member.", "warning")
                else:
                    db.session.add(TeamMember(team_id=team.id, user_id=member.id, created_at=datetime.now().isoformat(timespec="seconds")))
                    db.session.add(Notification(
                        user_id=member.id,
                        message=f"team:{team.id}: You were added to {team.name}.",
                        created_at=datetime.now().isoformat(timespec="seconds"),
                    ))
                    db.session.commit()
                    flash("Team member added.", "success")
            return redirect(url_for("teams"))
        name = request.form.get("name", "").strip()
        leader_id = request.form.get("leader_id", type=int)
        meeting_url = request.form.get("meeting_url", "").strip() or None
        requested_meeting_status = request.form.get("meeting_status", "Scheduled").strip()
        meeting_status = requested_meeting_status if requested_meeting_status in {"Scheduled", "Paused"} else "Scheduled"
        meeting_time = request.form.get("meeting_time", "").strip() or None
        if not name or not db.session.get(User, leader_id):
            flash("Team name and a valid leader are required.", "danger")
        else:
            team = Team(
                name=name,
                description=request.form.get("description", "").strip(),
                leader_id=leader_id,
                meeting_url=meeting_url,
                meeting_status=meeting_status,
                meeting_time=meeting_time,
                created_at=datetime.now().isoformat(timespec="seconds"),
            )
            db.session.add(team)
            try:
                db.session.flush()
                db.session.add(TeamMember(
                    team_id=team.id,
                    user_id=leader_id,
                    permissions="view_tasks,create_tasks,post_messages,manage_members,manage_meetings",
                    created_at=datetime.now().isoformat(timespec="seconds"),
                ))
                db.session.commit()
                flash("Team created successfully.", "success")
            except Exception:
                db.session.rollback()
                flash("A team with this name already exists.", "danger")
        return redirect(url_for("teams"))
    current_view = request.args.get("view", "all").strip().lower()
    teams_list = Team.query.order_by(Team.name).all()
    current_user_id = session.get("user_id")
    is_global_manager = session.get("role") in {"Admin", "Manager"}
    if current_view == "my" or not is_global_manager:
        current_team_ids = {
            team_id for (team_id,) in db.session.query(TeamMember.team_id).filter_by(user_id=current_user_id).all()
        }
        teams_list = [
            team for team in teams_list
            if team.leader_id == current_user_id or team.id in current_team_ids
        ]
    members = {team.id: TeamMember.query.filter_by(team_id=team.id).all() for team in teams_list}
    team_summary = {}
    for team in teams_list:
        team_tasks = Task.query.filter_by(team_id=team.id).all()
        overdue_count = sum(
            bool(task.due_date and task.status != "Completed" and task.due_date < date.today().isoformat())
            for task in team_tasks
        )
        team_summary[team.id] = {
            "total": len(team_tasks),
            "open": sum(task.status != "Completed" for task in team_tasks),
            "completed": sum(task.status == "Completed" for task in team_tasks),
            "overdue": overdue_count,
            "progress": round((sum(task.status == "Completed" for task in team_tasks) / len(team_tasks) * 100), 1) if team_tasks else 100.0,
        }
    notifications, unread_count = get_shared_data()
    return render_template(
        "teams.html",
        teams=teams_list,
        members=members,
        team_summary=team_summary,
        users=User.query.order_by(User.username).all(),
        notifications=notifications,
        unread_count=unread_count,
        current_view=current_view,
    )


@app.route("/team-workspace")
@login_required
def open_team_workspace():
    user_id = session["user_id"]
    role = session.get("role")
    team = db.session.get(Team, session.get("team_workspace_id"))

    if team and role not in {"Admin", "Manager"}:
        membership = TeamMember.query.filter_by(team_id=team.id, user_id=user_id).first()
        if team.leader_id != user_id and not membership:
            team = None

    if not team and role in {"Admin", "Manager"}:
        team = Team.query.order_by(Team.name).first()
    elif not team:
        membership = TeamMember.query.filter_by(user_id=user_id).order_by(TeamMember.id.asc()).first()
        team = db.session.get(Team, membership.team_id) if membership else Team.query.filter_by(leader_id=user_id).first()

    if not team:
        flash("No team workspace is available yet.", "info")
        return redirect(url_for("teams"))

    return redirect(url_for("team_detail", team_id=team.id))


@app.route("/teams/<int:team_id>/meeting/start", methods=["POST"])
@login_required
def start_meeting(team_id):
    team = db.session.get(Team, team_id)
    if not team:
        return redirect(url_for("teams"))
    
    member_row = TeamMember.query.filter_by(team_id=team.id, user_id=session["user_id"]).first()
    is_manager = session.get("role") in {"Admin", "Manager"}
    member_permissions = set((member_row.permissions or "").split(",")) if member_row else set()
    can_manage = is_manager or team.leader_id == session["user_id"] or "manage_meetings" in member_permissions

    if not can_manage:
        flash("You don't have permission to start a meeting.", "danger")
        return redirect(url_for("team_detail", team_id=team.id))

    if not team.meeting_url:
        flash("Please set a meeting URL in the workspace settings first.", "warning")
        return redirect(url_for("team_detail", team_id=team.id))
    if TeamMeeting.query.filter_by(team_id=team.id, ended_at=None).first():
        flash("A meeting is already active for this team.", "warning")
        return redirect(url_for("team_detail", team_id=team.id))

    team.meeting_status = "Live"
    meeting = TeamMeeting(
        team_id=team.id,
        started_at=datetime.now().isoformat(timespec="seconds"),
        started_by=session["user_id"]
    )
    db.session.add(meeting)
    db.session.flush()
    notify_team_live_meeting(team, meeting)
    
    db.session.commit()
    flash("Meeting is now LIVE. Members have been notified.", "success")
    return redirect(url_for("team_detail", team_id=team.id))


@app.route("/teams/<int:team_id>/meeting/end", methods=["POST"])
@login_required
def end_meeting(team_id):
    team = db.session.get(Team, team_id)
    if not team:
        return redirect(url_for("teams"))
    if not valid_csrf():
        flash("Invalid request. Please refresh and try again.", "danger")
        return redirect(url_for("team_detail", team_id=team.id))
    member_row = TeamMember.query.filter_by(team_id=team.id, user_id=session["user_id"]).first()
    member_permissions = set((member_row.permissions or "").split(",")) if member_row else set()
    can_manage = (
        session.get("role") in {"Admin", "Manager"}
        or team.leader_id == session["user_id"]
        or "manage_meetings" in member_permissions
    )
    if not can_manage:
        flash("You don't have permission to end this meeting.", "danger")
        return redirect(url_for("team_detail", team_id=team.id))
    meeting = TeamMeeting.query.filter_by(team_id=team_id, ended_at=None).order_by(TeamMeeting.id.desc()).first()
    if not meeting:
        flash("There is no active meeting to end.", "warning")
        return redirect(url_for("team_detail", team_id=team_id))

    end_time = datetime.now()
    start_time = datetime.fromisoformat(meeting.started_at)
    duration = int((end_time - start_time).total_seconds() / 60)

    meeting.ended_at = end_time.isoformat(timespec="seconds")
    meeting.duration_minutes = max(1, duration)
    
    team.meeting_status = "Scheduled"
    team.total_meetings_count = (team.total_meetings_count or 0) + 1
    team.total_meeting_minutes = (team.total_meeting_minutes or 0) + meeting.duration_minutes
    
    db.session.commit()
    flash(f"Meeting ended. Duration: {meeting.duration_minutes} minutes.", "info")
    return redirect(url_for("team_detail", team_id=team.id))


@app.route("/teams/<int:team_id>", methods=["GET", "POST"])

@login_required
def team_detail(team_id):
    team = db.session.get(Team, team_id)
    if not team:
        flash("Team not found.", "danger")
        return redirect(url_for("teams"))
    member_rows = TeamMember.query.filter_by(team_id=team.id).order_by(TeamMember.id.asc()).all()
    member_users = [db.session.get(User, row.user_id) for row in member_rows if db.session.get(User, row.user_id)]
    member_ids = [user.id for user in member_users]
    is_global_manager = session.get("role") in {"Admin", "Manager"}
    is_team_manager = is_global_manager or team.leader_id == session["user_id"]
    current_member_row = next((row for row in member_rows if row.user_id == session["user_id"]), None)
    current_permissions = set((current_member_row.permissions or "").split(",")) if current_member_row else set()
    can_manage_members = is_team_manager or "manage_members" in current_permissions
    can_manage_meetings = is_team_manager or "manage_meetings" in current_permissions
    can_create_tasks = is_team_manager or "create_tasks" in current_permissions
    can_post_messages = is_team_manager or "post_messages" in current_permissions
    is_async_request = request.headers.get("X-Requested-With") == "XMLHttpRequest"
    if not is_team_manager and (session["user_id"] not in member_ids or "view_tasks" not in current_permissions):
        flash("You do not have access to this team.", "danger")
        return redirect(url_for("teams"))

    session["team_workspace_id"] = team.id

    if request.method == "POST":
        if not valid_csrf():
            if is_async_request:
                return {"ok": False, "error": "Invalid request. Please refresh and try again."}, 400
            flash("Invalid request. Please refresh and try again.", "danger")
            return redirect(url_for("team_detail", team_id=team.id))

        action = request.form.get("action", "message")
        if action in {"add_member", "remove_member"}:
            if not can_manage_members:
                if is_async_request:
                    return {"ok": False, "error": "You do not have permission to manage team members."}, 403
                flash("You do not have permission to manage team members.", "danger")
                return redirect(url_for("team_detail", team_id=team.id))
            if action == "add_member":
                raw_user_ids = request.form.get("user_ids", request.form.get("user_id", "")).strip()
                id_parts = [part.strip() for part in raw_user_ids.split(",")]
                if not raw_user_ids or any(not part.isdecimal() or int(part) < 1 for part in id_parts):
                    error = "Enter one or more valid user IDs separated by commas."
                    if is_async_request:
                        return {"ok": False, "error": error}, 400
                    flash(error, "danger")
                    return redirect(url_for("team_detail", team_id=team.id))

                requested_ids = list(dict.fromkeys(int(part) for part in id_parts))
                found_users = User.query.filter(User.id.in_(requested_ids)).all()
                users_by_id = {user.id: user for user in found_users}
                missing_ids = [user_id for user_id in requested_ids if user_id not in users_by_id]
                if missing_ids:
                    error = "No user was found with ID(s): " + ", ".join(map(str, missing_ids))
                    if is_async_request:
                        return {"ok": False, "error": error}, 404
                    flash(error, "danger")
                    return redirect(url_for("team_detail", team_id=team.id))

                existing_ids = {
                    user_id for (user_id,) in db.session.query(TeamMember.user_id)
                    .filter(TeamMember.team_id == team.id, TeamMember.user_id.in_(requested_ids)).all()
                }
                if existing_ids:
                    error = "Already in this team: " + ", ".join(map(str, sorted(existing_ids)))
                    if is_async_request:
                        return {"ok": False, "error": error}, 409
                    flash(error, "warning")
                    return redirect(url_for("team_detail", team_id=team.id))

                added_users = [users_by_id[user_id] for user_id in requested_ids]
                for member_user in added_users:
                    db.session.add(TeamMember(
                        team_id=team.id,
                        user_id=member_user.id,
                        created_at=datetime.now().isoformat(timespec="seconds"),
                    ))
                    db.session.add(Notification(
                        user_id=member_user.id,
                        message=f"team:{team.id}: You were added to {team.name}.",
                        created_at=datetime.now().isoformat(timespec="seconds"),
                    ))
                db.session.commit()
                if is_async_request:
                    return {
                        "ok": True,
                        "members": [{"id": user.id, "username": user.username, "role": user.role} for user in added_users],
                    }, 201
                flash(f"Added {len(added_users)} team member(s).", "success")
                return redirect(url_for("team_detail", team_id=team.id))

            member_field = "user_id" if action == "add_member" else "member_id"
            member_id = request.form.get(member_field, type=int)
            member_user = db.session.get(User, member_id)
            member_row = TeamMember.query.filter_by(team_id=team.id, user_id=member_id).first()
            if not member_row:
                flash("This user is not a member of the team.", "warning")
            elif member_id == team.leader_id:
                flash("The team leader cannot be removed from the team.", "danger")
            else:
                db.session.delete(member_row)
                db.session.add(Notification(
                    user_id=member_id,
                    message=f"team:{team.id}: You were removed from {team.name}.",
                    created_at=datetime.now().isoformat(timespec="seconds"),
                ))
                db.session.commit()
                flash("Team member removed.", "success")
            return redirect(url_for("team_detail", team_id=team.id))
        if action == "create_team_task" and not can_create_tasks:
            flash("You do not have permission to create team tasks.", "danger")
            return redirect(url_for("team_detail", team_id=team.id))
        if action == "permissions":
            if not is_team_manager:
                flash("Only the team leader, managers, and admins can update team permissions.", "danger")
                return redirect(url_for("team_detail", team_id=team.id))
            member_id = request.form.get("member_id", type=int)
            member_row = TeamMember.query.filter_by(team_id=team.id, user_id=member_id).first()
            member_user = db.session.get(User, member_id)
            if not member_row or not member_user or member_user.role == "Admin":
                flash("Only regular team members can receive team permissions.", "danger")
                return redirect(url_for("team_detail", team_id=team.id))
            selected_permissions = [
                permission for permission, _label in TEAM_PERMISSION_OPTIONS
                if request.form.get(f"permission_{permission}") == "on"
            ]
            member_row.permissions = ",".join(selected_permissions) or "view_tasks"
            db.session.commit()
            flash(f"Permissions updated for {member_user.username}.", "success")
            return redirect(url_for("team_detail", team_id=team.id))
        if action == "update_team":
            if not is_team_manager:
                flash("Only the team leader, managers, and admins can update the team workspace.", "danger")
                return redirect(url_for("team_detail", team_id=team.id))
            new_name = request.form.get("name", team.name).strip()
            leader_id = request.form.get("leader_id", type=int) or team.leader_id
            if not new_name or not db.session.get(User, leader_id):
                flash("Team name and a valid leader are required.", "danger")
                return redirect(url_for("team_detail", team_id=team.id))
            previous_leader_id = team.leader_id
            team.name = new_name
            team.leader_id = leader_id
            team.description = request.form.get("description", "").strip()
            active_meeting = TeamMeeting.query.filter_by(team_id=team.id, ended_at=None).first()
            requested_meeting_status = request.form.get("meeting_status", "Scheduled").strip()
            team.meeting_status = "Live" if active_meeting else requested_meeting_status if requested_meeting_status in {"Scheduled", "Paused"} else "Scheduled"
            team.meeting_time = request.form.get("meeting_time", "").strip() or None
            team.meeting_url = request.form.get("meeting_url", "").strip() or None
            try:
                new_leader_row = TeamMember.query.filter_by(team_id=team.id, user_id=leader_id).first()
                if not new_leader_row:
                    db.session.add(TeamMember(
                        team_id=team.id,
                        user_id=leader_id,
                        permissions="view_tasks,create_tasks,post_messages,manage_members,manage_meetings",
                        created_at=datetime.now().isoformat(timespec="seconds"),
                    ))
                if previous_leader_id != leader_id:
                    previous_leader = TeamMember.query.filter_by(team_id=team.id, user_id=previous_leader_id).first()
                    if previous_leader:
                        permissions = set((previous_leader.permissions or "").split(","))
                        permissions.difference_update({"manage_members", "manage_meetings"})
                        previous_leader.permissions = ",".join(sorted(permissions)) or "view_tasks"
                db.session.commit()
                flash("Team workspace updated successfully.", "success")
            except Exception:
                db.session.rollback()
                flash("A team with this name already exists.", "danger")
            return redirect(url_for("team_detail", team_id=team.id))

        if action == "create_team_task":
            title = request.form.get("title", "").strip()
            assignee_id = request.form.get("assignee_id", type=int) or session["user_id"]
            priority = request.form.get("priority", "Medium").strip()
            due_date = request.form.get("due_date", "").strip() or None
            if due_date:
                try:
                    datetime.strptime(due_date, "%Y-%m-%d")
                except ValueError:
                    flash("Due date is invalid.", "danger")
                    return redirect(url_for("team_detail", team_id=team.id))
            if not title or priority not in VALID_PRIORITIES:
                flash("A valid task title and priority are required.", "danger")
                return redirect(url_for("team_detail", team_id=team.id))
            if assignee_id not in member_ids and session.get("role") not in {"Admin", "Manager"}:
                flash("You can only assign tasks to members of this team.", "danger")
                return redirect(url_for("team_detail", team_id=team.id))
            if assignee_id not in member_ids:
                flash("The assignee must be a member of this team.", "danger")
                return redirect(url_for("team_detail", team_id=team.id))
            task = Task(
                title=title,
                status="Pending",
                priority=priority,
                due_date=due_date,
                user_id=assignee_id,
                creator_id=session["user_id"],
                team_id=team.id,
            )
            db.session.add(task)
            db.session.flush()
            db.session.add(Detail(task_id=task.id, description=request.form.get("description", "").strip(), updated_at=datetime.now().isoformat(timespec="minutes")))
            if assignee_id != session["user_id"]:
                db.session.add(Notification(user_id=assignee_id, message=f"team_task:{task.id}: New team task assigned: {title}", created_at=datetime.now().isoformat(timespec="minutes")))
            write_audit("created", "task", task.id, new_value={"team_id": team.id, "title": title, "assignee": assignee_id})
            db.session.commit()
            record_realtime_event("TASK_CREATED", session["user_id"], "task", task.id, {"title": task.title, "status": task.status, "priority": task.priority, "teamId": team.id}, {f"team:{team.id}", f"user:{assignee_id}", f"user:{session['user_id']}"})
            flash("Team task created successfully.", "success")
            return redirect(url_for("team_detail", team_id=team.id))

        body = request.form.get("message", "").strip()
        if body and not can_post_messages:
            if is_async_request:
                return {"ok": False, "error": "You do not have permission to post team messages."}, 403
            flash("You do not have permission to post team messages.", "danger")
            return redirect(url_for("team_detail", team_id=team.id))
        if body:
            sender = db.session.get(User, session["user_id"])
            team_message = TeamMessage(team_id=team.id, user_id=session["user_id"], body=body, created_at=datetime.now().isoformat(timespec="seconds"))
            db.session.add(team_message)
            db.session.flush()
            for member in member_users:
                if member.id == session["user_id"]:
                    continue
                preview = body[:80] + ("..." if len(body) > 80 else "")
                db.session.add(Notification(
                    user_id=member.id,
                    message=f"team:{team.id}: {sender.username if sender else 'Someone'} posted in {team.name}: {preview}",
                    created_at=datetime.now().isoformat(timespec="seconds"),
                ))
            db.session.commit()
            record_realtime_event("COMMENT_CREATED", session["user_id"], "team", team.id, {"messageId": team_message.id, "body": team_message.body, "username": sender.username if sender else "User"}, {f"team:{team.id}"})
            if is_async_request:
                return {
                    "ok": True,
                    "message": {
                        "id": team_message.id,
                        "username": sender.username if sender else "User",
                        "body": team_message.body,
                        "created_at": team_message.created_at,
                    },
                }, 201
            flash("Message sent to the team.", "success")
        else:
            if is_async_request:
                return {"ok": False, "error": "Message cannot be empty."}, 400
            flash("Message cannot be empty.", "danger")
        return redirect(url_for("team_detail", team_id=team.id))

    team_tasks = Task.query.filter_by(team_id=team.id).order_by(Task.id.desc()).all()
    total_tasks = len(team_tasks)
    completed_tasks = sum(task.status == "Completed" for task in team_tasks)
    pending_tasks = sum(task.status != "Completed" for task in team_tasks)
    
    # New Stats for Team
    member_count = len(member_users)
    total_meetings = team.total_meetings_count or 0
    total_hours = round((team.total_meeting_minutes or 0) / 60, 1)
    
    overdue_tasks = sum(bool(task.due_date and task.status != "Completed" and task.due_date < date.today().isoformat()) for task in team_tasks)

    status_summary = {
        "Pending": sum(task.status == "Pending" for task in team_tasks),
        "In Progress": sum(task.status == "In Progress" for task in team_tasks),
        "Completed": completed_tasks,
    }
    priority_summary = {
        "High": sum(task.priority == "High" for task in team_tasks),
        "Medium": sum(task.priority == "Medium" for task in team_tasks),
        "Low": sum(task.priority == "Low" for task in team_tasks),
    }
    upcoming_tasks = sorted(
        [task for task in team_tasks if task.due_date and task.status != "Completed"],
        key=lambda task: task.due_date,
    )
    next_due_task = upcoming_tasks[0] if upcoming_tasks else None
    team_messages = TeamMessage.query.filter_by(team_id=team.id).order_by(TeamMessage.id.asc()).all()
    member_access = {
        row.user_id: set((row.permissions or "view_tasks,create_tasks,post_messages").split(","))
        for row in member_rows
    }
    member_stats = []
    for member in member_users:
        member_tasks = Task.query.filter_by(team_id=team.id, user_id=member.id).all()
        completed = sum(task.status == "Completed" for task in member_tasks)
        overdue = sum(bool(task.due_date and task.status != "Completed" and task.due_date < date.today().isoformat()) for task in member_tasks)
        member_stats.append({
            "user": member,
            "total": len(member_tasks),
            "completed": completed,
            "overdue": overdue,
            "open": sum(task.status != "Completed" for task in member_tasks),
            "health": round((completed / len(member_tasks) * 100), 1) if member_tasks else 100.0,
        })
    member_stats = sorted(member_stats, key=lambda item: (-item["health"], -item["total"]))
    notifications, unread_count = get_shared_data()
    user_lookup = {user.id: user for user in User.query.order_by(User.username).all()}
    return render_template(
        "team_detail.html",
        team=team,
        members=member_users,
        tasks=team_tasks,
        total_tasks=total_tasks,
        completed_tasks=completed_tasks,
        pending_tasks=pending_tasks,
        overdue_tasks=overdue_tasks,
        status_summary=status_summary,
        priority_summary=priority_summary,
        next_due_task=next_due_task,
                team_messages=team_messages,
        member_stats=member_stats,
        member_count=member_count,
        total_meetings=total_meetings,
        total_hours=total_hours,
        notifications=notifications,

        unread_count=unread_count,
        users=User.query.order_by(User.username).all(),
        member_access=member_access,
        permission_options=TEAM_PERMISSION_OPTIONS,
        user_lookup=user_lookup,
        member_ids=member_ids,
        is_team_manager=is_team_manager,
        can_manage_members=can_manage_members,
        can_manage_meetings=can_manage_meetings,
        can_create_tasks=can_create_tasks,
        can_post_messages=can_post_messages,
    )


@app.route("/roles", methods=["GET", "POST"])
@login_required
def roles():
    if not has_permission(session["user_id"], "roles.view"):
        flash("You do not have permission to view roles.", "danger")
        return redirect(url_for("tasks"))
    if request.method == "POST":
        submitted_actions = request.form.getlist("action")
        action = "delete" if "delete" in submitted_actions else submitted_actions[0] if submitted_actions else "save"
        role = db.session.get(Role, request.form.get("role_id", type=int)) if request.form.get("role_id") else None
        required = "roles.create" if action == "create" else "roles.delete" if action == "delete" else "roles.edit"
        if not has_permission(session["user_id"], required):
            return {"ok": False, "error": "You are not authorized for this role action."}, 403
        if action == "create":
            name = request.form.get("name", "").strip()
            if not name or len(name) > 80 or Role.query.filter_by(name=name).first():
                flash("Role name is required and must be unique.", "danger")
            else:
                role = Role(name=name, description=request.form.get("description", "").strip(), is_system=False, created_at=datetime.now().isoformat(timespec="seconds"))
                db.session.add(role)
                db.session.flush()
                write_audit("ROLE_CREATED", "role", role.id, new_value={"name": role.name})
                db.session.commit()
                flash("Role created successfully.", "success")
            return redirect(url_for("roles"))
        if not role:
            return {"ok": False, "error": "Role not found."}, 404
        if action == "delete":
            if role.name == "Super Admin" or role.is_system:
                return {"ok": False, "error": "System roles cannot be deleted."}, 403
            UserRole.query.filter_by(role_id=role.id).delete()
            RolePermission.query.filter_by(role_id=role.id).delete()
            db.session.delete(role)
            write_audit("ROLE_DELETED", "role", role.id, old_value={"name": role.name})
            db.session.commit()
            flash("Role deleted successfully.", "success")
            return redirect(url_for("roles"))
        old_permissions = {row.permission_id: row.scope for row in RolePermission.query.filter_by(role_id=role.id).all()}
        RolePermission.query.filter_by(role_id=role.id).delete()
        selected = []
        for key in ALL_PERMISSIONS:
            if request.form.get(f"permission_{key}") == "on":
                scope = request.form.get(f"scope_{key}", "ANY")
                if scope not in PERMISSION_SCOPES:
                    continue
                permission = Permission.query.filter_by(key=key).first()
                if permission:
                    db.session.add(RolePermission(role_id=role.id, permission_id=permission.id, scope=scope))
                    selected.append({"permission": key, "scope": scope})
        role.description = request.form.get("description", role.description or "").strip()
        write_audit("ROLE_PERMISSION_UPDATED", "role", role.id, old_value=old_permissions, new_value=selected)
        db.session.commit()
        affected_user_ids = [user_id for (user_id,) in db.session.query(UserRole.user_id).filter_by(role_id=role.id).all()]
        for user_id in affected_user_ids:
            record_realtime_event("PERMISSIONS_UPDATED", session["user_id"], "user", user_id, {"roleId": role.id, "roleName": role.name}, {f"user:{user_id}"})
        flash("Role permissions updated successfully.", "success")
        return redirect(url_for("roles"))
    role_rows = []
    for role in Role.query.order_by(Role.name).all():
        role_rows.append({"role": role, "permissions": {row.key: rp.scope for rp, row in db.session.query(RolePermission, Permission).join(Permission, Permission.id == RolePermission.permission_id).filter(RolePermission.role_id == role.id).all()}})
    return render_template("roles.html", role_rows=role_rows, permission_catalog=PERMISSION_CATALOG, scopes=sorted(PERMISSION_SCOPES), notifications=get_shared_data()[0], unread_count=get_shared_data()[1])


@app.route("/api/me/permissions")
@login_required
def my_permissions():
    rows = db.session.query(Permission.key, Role.name, RolePermission.scope).join(RolePermission, RolePermission.permission_id == Permission.id).join(Role, Role.id == RolePermission.role_id).join(UserRole, UserRole.role_id == Role.id).filter(UserRole.user_id == session["user_id"]).all()
    return {"permissions": [{"key": key, "role": role, "scope": scope} for key, role, scope in rows]}


@app.route("/users", methods=["GET", "POST"])
@login_required
def users():
    view = request.args.get("view", "all")
    if not has_permission(session["user_id"], "users.view"):
        flash("You do not have permission to view users.", "danger")
        return redirect(url_for("tasks"))
    if request.method == "POST":
        if not valid_csrf():
            flash("Invalid request. Please refresh and try again.", "danger")
            return redirect(url_for("users", view=view))
        action = request.form.get("action", "save")
        user_id = request.form.get("user_id", type=int)
        user = User.query.get(user_id) if user_id else None
        required_permission = "users.delete" if action == "delete" else "users.edit" if user else "users.create"
        if not has_permission(session["user_id"], required_permission):
            flash("You do not have permission to manage users this way.", "danger")
            return redirect(url_for("users", view=view))
        if action == "delete":
            if not user:
                flash("User not found.", "danger")
            elif user.id == session["user_id"]:
                flash("You cannot delete your own account.", "danger")
            elif is_super_admin(user.id) and not is_super_admin(session["user_id"]):
                flash("Only a Super Admin can delete a Super Admin account.", "danger")
            else:
                Task.query.filter_by(user_id=user.id).update({"user_id": None})
                Notification.query.filter_by(user_id=user.id).delete()
                db.session.delete(user)
                db.session.commit()
                flash("User deleted successfully.", "success")
            return redirect(url_for("users", view=view))
        if user:
            user.first_name = request.form.get("first_name", user.first_name).strip()
            user.last_name = request.form.get("last_name", user.last_name).strip()
            user.phone = request.form.get("phone", user.phone)
            user.email = request.form.get("email", user.email)
            role = request.form.get("role")
            if role in VALID_ROLES and user.id != session["user_id"]:
                user.role = role
            new_password = request.form.get("new_password", "")
            if new_password:
                if len(new_password) < 8:
                    flash("Password must contain at least 8 characters.", "danger")
                    return redirect(url_for("users", view=view))
                user.password_hash = generate_password_hash(new_password)
            submitted_role_ids = [int(value) for value in request.form.getlist("role_ids") if value.isdigit()]
            if not submitted_role_ids:
                legacy_role = Role.query.filter_by(name=user.role).first()
                submitted_role_ids = [legacy_role.id] if legacy_role else []
            roles_ok, roles_error = sync_user_roles(user, submitted_role_ids, session["user_id"])
            if not roles_ok:
                db.session.rollback()
                flash(roles_error, "danger")
                return redirect(url_for("users", view=view))
            db.session.commit()
            record_realtime_event("PERMISSIONS_UPDATED", session["user_id"], "user", user.id, {"reason": "roles_changed"}, {f"user:{user.id}"})
            flash("User updated successfully.", "success")
            return redirect(url_for("users", view=view))
        first_name = request.form.get("first_name", "").strip()
        last_name = request.form.get("last_name", "").strip()
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        role = request.form.get("role", "User")
        if not first_name or not last_name or not username or len(password) < 8:
            flash("Names, username, and an 8-character password are required.", "danger")
            return redirect(url_for("users", view=view))
        if role not in VALID_ROLES or User.query.filter_by(username=username).first():
            flash("Username already exists or role is invalid.", "danger")
            return redirect(url_for("users", view=view))
        new_user = User(first_name=first_name, last_name=last_name, username=username,
                            phone=request.form.get("phone", "").strip(), email=request.form.get("email", "").strip() or None, gender=request.form.get("gender", "").strip(),
                            password_hash=generate_password_hash(password), role=role,
                            created_at=datetime.now().isoformat(timespec="minutes"))
        db.session.add(new_user)
        db.session.flush()
        submitted_role_ids = [int(value) for value in request.form.getlist("role_ids") if value.isdigit()]
        if not submitted_role_ids:
            default_role = Role.query.filter_by(name=role if role in VALID_ROLES else "User").first()
            submitted_role_ids = [default_role.id] if default_role else []
        roles_ok, roles_error = sync_user_roles(new_user, submitted_role_ids, session["user_id"])
        if not roles_ok:
            db.session.rollback()
            flash(roles_error, "danger")
            return redirect(url_for("users", view=view))
        db.session.commit()
        flash("User added successfully.", "success")
        return redirect(url_for("users", view=view))

    users_query = db.session.execute(text("SELECT id, first_name, last_name, username, phone, gender, NULL AS password, role, created_at, email FROM users ORDER BY id ASC")).fetchall()
    if view == "managers":
        users_query = [user for user in users_query if user[7] == "Manager"]
    elif view == "admins":
        users_query = [user for user in users_query if user[7] == "Admin"]
    notifications, unread_count = get_shared_data()
    role_assignments = {user.id: {role.name for role in user_roles(user)} for user in User.query.all()}
    return render_template("users.html", users_list=users_query, current_view=view, roles=Role.query.order_by(Role.name).all(), role_assignments=role_assignments, notifications=notifications, unread_count=unread_count)


@app.route("/tasks")
@login_required
def tasks():
    if not has_permission(session["user_id"], "tasks.view"):
        flash("You do not have permission to view tasks.", "danger")
        return redirect(url_for("index"))
    tasks_query = task_rows_for_current_user()
    requested_view = request.args.get("view", "all")
    initial_tab = {"all": "all", "sent": "sent", "received": "received", "urgent": "urgent"}.get(requested_view, "inbox")
    focus_task = None
    focus_id = request.args.get("focus", type=int)
    if focus_id:
        focused = db.session.get(Task, focus_id)
        member_team_ids = {team_id for (team_id,) in db.session.query(TeamMember.team_id).filter_by(user_id=session["user_id"]).all()}
        if focused and (session.get("role") == "Admin" or focused.user_id == session["user_id"] or focused.team_id in member_team_ids):
            detail = Detail.query.filter_by(task_id=focused.id).first()
            assignee = db.session.get(User, focused.user_id) if focused.user_id else None
            focus_task = {
                "id": focused.id,
                "title": focused.title or "",
                "assignee": assignee.username if assignee else "Unassigned",
                "assignee_id": focused.user_id or "",
                "sender_id": focused.creator_id or "",
                "can_change_status": session.get("role") == "Admin" or (focused.user_id == session["user_id"] and focused.creator_id != session["user_id"]),
                "can_manage_task": session.get("role") == "Admin" or focused.creator_id == session["user_id"],
                "priority": focused.priority or "Medium",
                "status": focused.status or "Pending",
                "due_date": focused.due_date or "",
                "description": detail.description if detail else "No description provided",
            }
    status_filter = request.args.get("status", "All")
    priority_filter = request.args.get("priority", "All")
    search = request.args.get("q", "").strip().lower()
    if status_filter in VALID_STATUSES:
        tasks_query = [task for task in tasks_query if task.status == status_filter]
    if priority_filter in VALID_PRIORITIES:
        tasks_query = [task for task in tasks_query if task.priority == priority_filter]
    if search:
        tasks_query = [task for task in tasks_query if search in (task.title or "").lower()]
    total_tasks = len(tasks_query)
    completed_tasks = sum(task.status == "Completed" for task in tasks_query)
    current_date = date.today().isoformat()
    overdue_tasks = sum(bool(task.due_date and task.status != "Completed" and task.due_date < current_date) for task in tasks_query)
    notifications, unread_count = get_shared_data()
    return render_template("tasks.html", tasks=tasks_query, users_list=User.query.order_by(User.username).all(),
                           total_tasks=total_tasks, completed_tasks=completed_tasks, pending_tasks=total_tasks - completed_tasks,
                           overdue_tasks=overdue_tasks, current_date_str=current_date, notifications=notifications,
                           unread_count=unread_count, focus_task=focus_task, initial_tab=initial_tab)


@app.route("/add_task", methods=["GET", "POST"])
@login_required
def add_task():
    if not has_permission(session["user_id"], "tasks.create"):
        flash("You do not have permission to create tasks.", "danger")
        return redirect(url_for("tasks"))
    if request.method == "GET":
        return render_template("add_task.html", users_list=User.query.order_by(User.username).all(), teams=Team.query.order_by(Team.name).all())
    if not valid_csrf():
        flash("Invalid request. Please refresh and try again.", "danger")
        return redirect(url_for("tasks"))
    title = request.form.get("title", "").strip()
    priority = request.form.get("priority", "Medium")
    due_date = request.form.get("due_date", "") or None
    if not title or len(title) > 100 or priority not in VALID_PRIORITIES:
        flash("A valid title and priority are required.", "danger")
        return redirect(url_for("tasks"))
    if due_date:
        try:
            datetime.strptime(due_date, "%Y-%m-%d")
        except ValueError:
            flash("Due date is invalid.", "danger")
            return redirect(url_for("tasks"))
    assigned_user_id = request.form.get("user_id", type=int)
    if not assigned_user_id or not User.query.get(assigned_user_id):
        flash("Enter a valid assignee user ID.", "danger")
        return redirect(url_for("tasks"))
    team_id = None
    if session.get("role") == "Admin":
        team_id = request.form.get("team_id", type=int)
        if team_id and not Team.query.get(team_id):
            flash("Assigned team was not found.", "danger")
            return redirect(url_for("tasks"))
    task = Task(title=title, status="Pending", priority=priority, due_date=due_date,
                user_id=assigned_user_id, creator_id=session["user_id"], team_id=team_id)
    db.session.add(task)
    db.session.flush()
    db.session.add(Detail(task_id=task.id, description=request.form.get("description", "").strip(), updated_at=datetime.now().isoformat(timespec="minutes")))
    if assigned_user_id != session["user_id"]:
        db.session.add(Notification(user_id=assigned_user_id, message=f"New task assigned: {title}", created_at=datetime.now().isoformat(timespec="minutes")))
    write_audit("created", "task", task.id, new_value={"title": title, "assignee": assigned_user_id})
    db.session.commit()
    assigned_user = db.session.get(User, assigned_user_id)
    preference = "task_created" if assigned_user_id == session["user_id"] else "task_assigned"
    send_email_notification(assigned_user.email if assigned_user else None, f"New task assigned: {title}", f"You have been assigned a new task: {title}", assigned_user_id, preference)
    db.session.commit()
    record_realtime_event("TASK_CREATED", session["user_id"], "task", task.id, {"title": task.title, "status": task.status, "priority": task.priority, "dueDate": task.due_date}, {f"user:{assigned_user_id}", f"user:{session['user_id']}"})
    record_realtime_event("NOTIFICATION_CREATED", session["user_id"], "notification", assigned_user_id, {"message": f"New task assigned: {title}"}, {f"user:{assigned_user_id}"})
    flash("Task created successfully.", "success")
    return redirect(url_for("tasks"))


@app.route("/update_task_status/<int:id>/<string:new_status>", methods=["POST"])
@login_required
def update_status(id, new_status):
    if not valid_csrf():
        flash("Invalid request. Please refresh and try again.", "danger")
        return redirect(url_for("tasks"))
    if new_status not in VALID_STATUSES:
        flash("Invalid status.", "danger")
        return redirect(url_for("tasks"))
    task = db.session.get(Task, id)
    if not task or not has_permission(session["user_id"], "tasks.change_status", task):
        flash("You do not have permission to update this task.", "danger")
        return redirect(url_for("tasks"))
    comment_body = request.form.get("comment", "").strip()
    if new_status in {"Completed", "Rejected"} and not comment_body:
        flash("Add a comment before completing or rejecting this task.", "danger")
        return redirect(url_for("tasks", focus=task.id))
    if new_status == "Completed" and session.get("role") != "Admin":
        blockers = [dependency for dependency in dependency_predecessors(id) if dependency.status != "Completed"]
        if blockers:
            flash("This task is blocked until its prerequisite tasks are completed.", "danger")
            return redirect(url_for("tasks"))
    old_status = task.status
    task.status = new_status
    task.completed_at = datetime.now().isoformat(timespec="seconds") if new_status == "Completed" else None
    if comment_body:
        db.session.add(Comment(task_id=task.id, user_id=session["user_id"], body=comment_body, created_at=datetime.now().isoformat(timespec="seconds")))
    db.session.commit()
    creator = db.session.get(User, task.creator_id) if task.creator_id else None
    if creator and creator.id != session["user_id"]:
        send_email_notification(creator.email, f"Task status changed: {task.title}", f"Task '{task.title}' is now {new_status}.", creator.id, "status_changed")
        db.session.commit()
    write_audit("status_changed", "task", task.id, old_value={"status": old_status}, new_value={"status": new_status})
    db.session.commit()
    record_realtime_event("TASK_STATUS_CHANGED", session["user_id"], "task", task.id, {"oldStatus": old_status, "newStatus": new_status}, {f"task:{task.id}", f"user:{task.user_id}", f"user:{task.creator_id}"})
    if creator and creator.id != session["user_id"]:
        record_realtime_event("NOTIFICATION_CREATED", session["user_id"], "notification", creator.id, {"message": f"Task status changed: {task.title}"}, {f"user:{creator.id}"})
    flash("Task status updated successfully.", "success")
    return redirect(url_for("tasks"))


@app.route("/tasks/<int:id>/dependencies", methods=["GET", "POST"])
@login_required
def task_dependencies(id):
    task = db.session.get(Task, id)
    if not task:
        flash("Task not found.", "danger")
        return redirect(url_for("tasks"))
    if request.method == "POST":
        if session.get("role") != "Admin" and task.creator_id != session["user_id"]:
            flash("Only the task sender can manage dependencies.", "danger")
            return redirect(url_for("tasks"))
        if not valid_csrf():
            flash("Invalid request. Please try again.", "danger")
            return redirect(url_for("task_dependencies", id=id))
        predecessor_id = request.form.get("predecessor_id", type=int)
        dependency_type = request.form.get("dependency_type", "Blocks")
        predecessor = db.session.get(Task, predecessor_id) if predecessor_id else None
        if not predecessor or predecessor.id == task.id or dependency_type not in VALID_DEPENDENCY_TYPES:
            flash("Choose a valid different task and dependency type.", "danger")
        elif dependency_would_cycle(predecessor.id, task.id):
            flash("That dependency would create a cycle.", "danger")
        else:
            db.session.add(TaskDependency(predecessor_id=predecessor.id, successor_id=task.id, dependency_type=dependency_type, created_by=session["user_id"], created_at=datetime.now().isoformat(timespec="seconds")))
            try:
                db.session.commit()
                flash("Dependency added successfully.", "success")
            except Exception:
                db.session.rollback()
                flash("This dependency already exists.", "danger")
        return redirect(url_for("task_dependencies", id=id))
    dependencies = db.session.query(TaskDependency, Task).join(Task, Task.id == TaskDependency.predecessor_id).filter(TaskDependency.successor_id == id).all()
    candidates = Task.query.filter(Task.id != id).order_by(Task.title).all()
    notifications, unread_count = get_shared_data()
    return render_template("dependencies.html", task=task, dependencies=dependencies, candidates=candidates, dependency_types=sorted(VALID_DEPENDENCY_TYPES), notifications=notifications, unread_count=unread_count)


@app.route("/tasks/<int:task_id>/dependencies/<int:dependency_id>/delete", methods=["POST"])
@login_required
def delete_dependency(task_id, dependency_id):
    task = db.session.get(Task, task_id)
    dependency = db.session.get(TaskDependency, dependency_id)
    if not task or not dependency or dependency.successor_id != task_id:
        flash("Dependency not found.", "danger")
    elif session.get("role") != "Admin" and task.creator_id != session["user_id"]:
        flash("Only the task sender can manage dependencies.", "danger")
    elif not valid_csrf():
        flash("Invalid request. Please try again.", "danger")
    else:
        db.session.delete(dependency)
        db.session.commit()
        flash("Dependency removed.", "success")
    return redirect(url_for("task_dependencies", id=task_id))


@app.route("/reports")
@login_required
def reports():
    if not has_permission(session["user_id"], "reports.view"):
        flash("You do not have permission to view reports.", "danger")
        return redirect(url_for("tasks"))
    reports_data = task_rows_for_current_user()
    notifications, unread_count = get_shared_data()
    return render_template("reports.html", reports=reports_data, users_count=User.query.count(), notifications=notifications, unread_count=unread_count)


@app.route("/analytics")
@login_required
def analytics():
    if not has_permission(session["user_id"], "reports.view"):
        flash("You do not have permission to view analytics.", "danger")
        return redirect(url_for("tasks"))
    visible = visible_tasks_for_user(session["user_id"])
    today = date.today().isoformat()
    status_counts = {status: sum(task.status == status for task in visible) for status in sorted(VALID_STATUSES)}
    priority_counts = {priority: sum(task.priority == priority for task in visible) for priority in sorted(VALID_PRIORITIES)}
    overdue = sum(bool(task.due_date and task.status != "Completed" and task.due_date < today) for task in visible)
    completed_by_day = {}
    for task in visible:
        if task.completed_at:
            day = task.completed_at[:10]
            completed_by_day[day] = completed_by_day.get(day, 0) + 1
    team_counts = {}
    for task in visible:
        if task.team_id:
            team = db.session.get(Team, task.team_id)
            name = team.name if team else "Unknown"
            team_counts[name] = team_counts.get(name, 0) + 1

    team_health = []
    all_teams = Team.query.order_by(Team.name).all()
    for team in all_teams:
        member_ids = [row.user_id for row in TeamMember.query.filter_by(team_id=team.id).all()]
        if session.get("role") not in {"Admin", "Manager"} and session["user_id"] not in member_ids:
            continue
        team_tasks = Task.query.filter((Task.team_id == team.id) | (Task.user_id.in_(member_ids))).all()
        completed = sum(task.status == "Completed" for task in team_tasks)
        overdue_tasks = sum(bool(task.due_date and task.status != "Completed" and task.due_date < today) for task in team_tasks)
        open_tasks = sum(task.status != "Completed" for task in team_tasks)
        health = round((completed / len(team_tasks) * 100), 1) if team_tasks else 100.0
        team_health.append({
            "name": team.name,
            "members": len(member_ids),
            "tasks": len(team_tasks),
            "completed": completed,
            "open": open_tasks,
            "overdue": overdue_tasks,
            "meeting_status": team.meeting_status or "Scheduled",
            "health": health,
        })

    team_health = sorted(team_health, key=lambda item: (-item["health"], -item["open"]))
    live_meetings = sum(1 for item in team_health if item["meeting_status"] == "Live")
    urgent_teams = sum(1 for item in team_health if item["overdue"] > 0)
    notifications, unread_count = get_shared_data()
    return render_template(
        "analytics.html",
        total=len(visible),
        status_counts=status_counts,
        priority_counts=priority_counts,
        overdue=overdue,
        completed_by_day=dict(sorted(completed_by_day.items())[-14:]),
        team_counts=team_counts,
        team_health=team_health,
        total_teams=len(team_health),
        live_meetings=live_meetings,
        urgent_teams=urgent_teams,
        notifications=notifications,
        unread_count=unread_count,
    )


@app.route("/audit-logs")
@login_required
def audit_logs():
    if not has_permission(session["user_id"], "audit_logs.view"):
        flash("You do not have permission to view audit logs.", "danger")
        return redirect(url_for("tasks"))
    logs = AuditLog.query.order_by(AuditLog.id.desc()).limit(500).all()
    notifications, unread_count = get_shared_data()
    return render_template("audit_logs.html", logs=logs, notifications=notifications, unread_count=unread_count)


@app.route("/tasks/<int:id>/comments", methods=["GET", "POST"])
@login_required
def task_comments(id):
    task = db.session.get(Task, id)
    if not task or not has_permission(session["user_id"], "comments.view", task) or not can_access_task(session["user_id"], task):
        flash("Task not found or not accessible.", "danger")
        return redirect(url_for("tasks"))
    if request.method == "POST":
        if not valid_csrf():
            flash("Invalid request. Please try again.", "danger")
        else:
            if not has_permission(session["user_id"], "comments.create", task):
                flash("You do not have permission to create comments.", "danger")
                return redirect(url_for("task_comments", id=id))
            body = request.form.get("body", "").strip()
            if body and len(body) <= 4000:
                comment = Comment(task_id=id, user_id=session["user_id"], body=body, created_at=datetime.now().isoformat(timespec="seconds"))
                db.session.add(comment)
                write_audit("created", "comment", comment.id, new_value={"task_id": id})
                mentions = set(re.findall(r"@([A-Za-z0-9_.-]+)", body))
                for username in mentions:
                    mentioned = User.query.filter(func.lower(User.username) == username.lower()).first()
                    if mentioned and mentioned.id != session["user_id"]:
                        db.session.add(Notification(user_id=mentioned.id, message=f"You were mentioned in a task comment: {task.title}", created_at=datetime.now().isoformat(timespec="seconds")))
                db.session.commit()
                record_realtime_event("COMMENT_CREATED", session["user_id"], "task", task.id, {"commentId": comment.id, "body": comment.body}, {f"task:{task.id}", f"user:{task.user_id}", f"user:{task.creator_id}"})
                flash("Comment added.", "success")
            else:
                flash("Comment must contain 1 to 4000 characters.", "danger")
    comments = Comment.query.filter_by(task_id=id).order_by(Comment.id.asc()).all()
    notifications, unread_count = get_shared_data()
    return render_template("comments.html", task=task, comments=comments, notifications=notifications, unread_count=unread_count)


@app.route("/automation", methods=["GET", "POST"])
@admin_required
def automation():
    if request.method == "POST":
        if not valid_csrf():
            flash("Invalid request. Please try again.", "danger")
        else:
            rule = AutomationRule(name=request.form.get("name", "").strip(), event=request.form.get("event", "status_changed"), condition=request.form.get("condition", "equals"), value=request.form.get("value", "").strip(), action=request.form.get("action", "notify"), created_by=session["user_id"], created_at=datetime.now().isoformat(timespec="seconds"))
            if not rule.name or not rule.value:
                flash("Rule name and value are required.", "danger")
            else:
                db.session.add(rule)
                write_audit("created", "automation_rule", None, new_value={"name": rule.name, "event": rule.event})
                db.session.commit()
                flash("Automation rule created.", "success")
        return redirect(url_for("automation"))
    rules = AutomationRule.query.order_by(AutomationRule.id.desc()).all()
    notifications, unread_count = get_shared_data()
    return render_template("automation.html", rules=rules, notifications=notifications, unread_count=unread_count)


@app.route("/automation/<int:id>/toggle", methods=["POST"])
@admin_required
def toggle_automation(id):
    if not valid_csrf():
        flash("Invalid request. Please try again.", "danger")
    else:
        rule = db.session.get(AutomationRule, id)
        if rule:
            rule.enabled = not rule.enabled
            write_audit("toggled", "automation_rule", rule.id, new_value={"enabled": rule.enabled})
            db.session.commit()
            flash("Automation rule updated.", "success")
    return redirect(url_for("automation"))


@app.route("/data", methods=["GET", "POST"])
@admin_required
def data_center():
    preview = None
    errors = []
    resource = request.form.get("resource", "tasks")
    if request.method == "POST":
        if not valid_csrf():
            flash("Invalid request. Please try again.", "danger")
        elif request.form.get("action") == "confirm":
            rows = session.pop("import_preview", [])
            for row in rows:
                if resource == "users":
                    db.session.add(User(first_name=row.get("first_name", "Imported"), last_name=row.get("last_name", "User"), username=row["username"], password_hash=generate_password_hash(row["password"]), email=row.get("email") or None, role=row.get("role", "User"), created_at=datetime.now().isoformat(timespec="seconds")))
                elif resource == "teams":
                    leader = User.query.filter(func.lower(User.username) == row["leader_username"].lower()).first()
                    if leader:
                        team = Team(name=row["name"], description=row.get("description", ""), leader_id=leader.id, created_at=datetime.now().isoformat(timespec="seconds"))
                        db.session.add(team)
                        db.session.flush()
                        db.session.add(TeamMember(
                            team_id=team.id,
                            user_id=leader.id,
                            permissions="view_tasks,create_tasks,post_messages,manage_members,manage_meetings",
                            created_at=datetime.now().isoformat(timespec="seconds"),
                        ))
                else:
                    creator = session["user_id"]
                    db.session.add(Task(title=row["title"], priority=row.get("priority", "Medium") if row.get("priority", "Medium") in VALID_PRIORITIES else "Medium", status=row.get("status", "Pending") if row.get("status", "Pending") in VALID_STATUSES else "Pending", due_date=row.get("due_date") or None, creator_id=creator, user_id=creator))
            write_audit("imported", resource, new_value={"rows": len(rows)})
            db.session.commit()
            flash(f"Imported {len(rows)} valid row(s).", "success")
            return redirect(url_for("data_center"))
        else:
            try:
                rows, file_type = parse_import_file(request.files.get("file"))
                valid, errors = validate_import_rows(resource, rows)
                session["import_preview"] = valid[:1000]
                preview = valid[:25]
                flash(f"Preview ready from {file_type.upper()}: {len(valid)} valid row(s), {len(errors)} error(s).", "success")
            except (ValueError, AttributeError) as error:
                flash(str(error), "danger")
    notifications, unread_count = get_shared_data()
    return render_template("data_center.html", resource=resource, preview=preview, errors=errors, notifications=notifications, unread_count=unread_count)


@app.route("/export/<resource>.<file_format>")
@admin_required
def export_data(resource, file_format):
    if resource not in {"tasks", "users", "teams", "reports"} or file_format not in {"csv", "xlsx", "pdf"}:
        return {"ok": False, "error": "Unsupported export"}, 400
    if resource in {"tasks", "reports"}:
        rows = [{"id": task.id, "title": task.title, "status": task.status, "priority": task.priority, "due_date": task.due_date, "user_id": task.user_id, "team_id": task.team_id} for task in Task.query.order_by(Task.id).all()]
    elif resource == "users":
        rows = [{"id": user.id, "username": user.username, "role": user.role, "email": user.email} for user in User.query.order_by(User.id).all()]
    else:
        rows = [{"id": team.id, "name": team.name, "leader_id": team.leader_id, "status": team.status} for team in Team.query.order_by(Team.id).all()]
    if file_format == "csv":
        stream = io.StringIO()
        writer = csv.DictWriter(stream, fieldnames=list(rows[0].keys()) if rows else ["id"])
        writer.writeheader(); writer.writerows(rows)
        return send_file(io.BytesIO(stream.getvalue().encode("utf-8-sig")), mimetype="text/csv", as_attachment=True, download_name=f"{resource}.csv")
    if file_format == "xlsx":
        try:
            from openpyxl import Workbook
        except ImportError:
            return {"ok": False, "error": "Install openpyxl for Excel export"}, 501
        workbook = Workbook(); sheet = workbook.active
        headers = list(rows[0].keys()) if rows else ["id"]; sheet.append(headers)
        for row in rows: sheet.append([row.get(header) for header in headers])
        output = io.BytesIO(); workbook.save(output); output.seek(0)
        return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name=f"{resource}.xlsx")
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
    except ImportError:
        return {"ok": False, "error": "Install reportlab for PDF export"}, 501
    output = io.BytesIO(); pdf = canvas.Canvas(output, pagesize=letter); y = 760
    pdf.setFont("Helvetica", 9); pdf.drawString(40, y, resource.title()); y -= 22
    for row in rows[:100]:
        pdf.drawString(40, y, " | ".join(f"{key}: {value}" for key, value in row.items())[:110]); y -= 14
        if y < 40: pdf.showPage(); y = 760
    pdf.save(); output.seek(0)
    return send_file(output, mimetype="application/pdf", as_attachment=True, download_name=f"{resource}.pdf")


@app.route("/edit_task/<int:id>", methods=["GET", "POST"])
@login_required
def edit_task(id):
    task = db.session.get(Task, id)
    if not task or not has_permission(session["user_id"], "tasks.edit", task):
        flash("Only the task sender can edit this task.", "danger")
        return redirect(url_for("tasks"))
    detail = Detail.query.filter_by(task_id=id).first()
    if request.method == "GET":
        return render_template("add_task.html", users_list=User.query.order_by(User.username).all(), task=task, detail=detail, editing=True)
    if not valid_csrf():
        flash("Invalid request. Please refresh and try again.", "danger")
        return redirect(url_for("tasks"))
    title = request.form.get("title", "").strip()
    priority = request.form.get("priority", "Medium")
    due_date = request.form.get("due_date", "") or None
    if not title or len(title) > 100 or priority not in VALID_PRIORITIES:
        flash("A valid title and priority are required.", "danger")
        return redirect(url_for("tasks"))
    task.title, task.priority, task.due_date = title, priority, due_date
    if detail:
        detail.description = request.form.get("description", "").strip()
        detail.updated_at = datetime.now().isoformat(timespec="minutes")
    else:
        db.session.add(Detail(task_id=id, description=request.form.get("description", "").strip(), updated_at=datetime.now().isoformat(timespec="minutes")))
    db.session.commit()
    record_realtime_event("TASK_UPDATED", session["user_id"], "task", task.id, {"title": task.title, "priority": task.priority, "dueDate": task.due_date}, {f"task:{task.id}", f"user:{task.user_id}", f"user:{task.creator_id}"})
    flash("Task updated successfully.", "success")
    return redirect(url_for("tasks"))


@app.route("/delete_task/<int:id>", methods=["POST"])
@login_required
def delete_task(id):
    if not valid_csrf():
        flash("Invalid request. Please try again.", "danger")
        return redirect(url_for("tasks"))
    task = db.session.get(Task, id)
    if not task or not has_permission(session["user_id"], "tasks.delete", task):
        flash("Only the task sender can recall this task.", "danger")
        return redirect(url_for("tasks"))
    delete_comment = request.form.get("comment", "").strip()
    if not delete_comment:
        flash("Add a reason before deleting this task.", "danger")
        return redirect(url_for("tasks", focus=task.id))
    Detail.query.filter_by(task_id=id).delete()
    deleted_title = task.title
    deleted_user_id = task.user_id
    deleted_creator_id = task.creator_id
    write_audit("deleted", "task", task.id, old_value={"title": task.title, "status": task.status, "comment": delete_comment})
    db.session.delete(task)
    db.session.commit()
    record_realtime_event("TASK_DELETED", session["user_id"], "task", id, {"title": deleted_title, "reason": delete_comment, "recipientIds": [deleted_user_id, deleted_creator_id]}, {f"user:{deleted_user_id}", f"user:{deleted_creator_id}"})
    flash("Task deleted successfully.", "success")
    return redirect(url_for("tasks"))


if __name__ == "__main__":
    socketio.run(app, debug=os.environ.get("FLASK_DEBUG") == "1")
